"""Compost Intensive slides for p20 (slide 6) and j20270110c, made from Linnea's own Compost Intensive deck
(09-11-2026-compost-intensive-january-cohort-readable.pptx): one slide kept, only the words changed, every layer
and font as she set it (Linnea, 31 Oct: "we already made better visuals for compost intensive").

  python3 tools/wormhole-visuals/ci_from_deck.py
"""
import copy
import os

from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
DECKS = os.path.join(HERE, "..", "..", "apps", "web", "public", "wormhole", "decks")
SRC = os.path.join(DECKS, "09-11-2026-compost-intensive-january-cohort-readable.pptx")


def set_lines(shape, lines):
    """New words, same formatting: each line takes the first run's font of the matching paragraph."""
    tf = shape.text_frame
    paras = tf.paragraphs
    while len(paras) < len(lines):
        tf._txBody.append(copy.deepcopy(paras[-1]._p))
        paras = tf.paragraphs
    for p, line in zip(paras, lines):
        runs = p.runs
        runs[0].text = line
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    for p in paras[len(lines):]:
        p._p.getparent().remove(p._p)


def keep_only(prs, index):
    ids = prs.slides._sldIdLst
    for i, sid in reversed(list(enumerate(list(ids)))):
        if i != index:
            prs.part.drop_rel(sid.rId)
            ids.remove(sid)


def build(index, texts, out):
    prs = Presentation(SRC)
    keep_only(prs, index)
    shapes = {s.name: s for s in prs.slides[0].shapes}
    for name, lines in texts.items():
        set_lines(shapes[name], lines)
    prs.save(os.path.join(DECKS, out))
    return out


if __name__ == "__main__":
    print(build(0, {"Headline": ["Last call:", "starts tomorrow"],
                    "Details": ["Compost Intensive, 11 January"]},
                "10-01-2027-compost-intensive-last-call-v7.pptx"))
    print(build(2, {"Headline 2": ["Starts 11 January"],
                    "Details": ["Compost Intensive Course,", "10 weeks online. Link in bio."]},
                "02-11-2026-inside-a-compost-pile-slide-6-v7.pptx"))
