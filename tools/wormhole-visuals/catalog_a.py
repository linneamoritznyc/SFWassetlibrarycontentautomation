"""Catalog swap, worker A (27 Oct): new versions of posts whose photos repeat an earlier post.

Every repeated photo is swapped for a fresh photo from Linnea's SFW image catalog. Tag words,
tag colour, tag position, logo and size stay the same as the current slide; only the photo
changes. Deck-built slides (Costa Rica stories and carousels) keep their deck layout: the new
photo is laid into the deck's photo box, under the same text, flowers and logo.

  BOX_SRC=<catalog photos> SFW_FONTS=<woff2 dir> python3 tools/wormhole-visuals/catalog_a.py
    -> apps/web/public/wormhole/tiles/<id>-c1-<n>.jpg  (never overwrites; bumps to c2, c3)
"""
import os
import subprocess
import sys
import tempfile

from PIL import Image, ImageOps

import boxes as B
import photo_first as P
from photo_first import BL, BR, D, G, T, pic

OUT = B.OUT
TILES = OUT
DECKS = os.path.join(B.REPO, "apps", "web", "public", "wormhole", "decks")
SFW = "Photo: Soil Food Web Foundation"
BREVITY = "Photo: Brevity Media"

P.CR.update({
    "R5A_4016.jpg": BREVITY,
    "R5A_3994.jpg": BREVITY,
    "20250617_135603.jpg": SFW,
    "IMG_4534.jpg": SFW,
    "20250625_155834.jpg": SFW,
    "20250626_093348.jpg": SFW,
    "20250619_071702.jpg": SFW,
    "20250620_153134.jpg": SFW,
    "20250620_144146.jpg": SFW,
    "20250618_111904.jpg": SFW,
    "IMG_2593.jpg": SFW,
    "20250311_112306.jpg": SFW,
    "PXL_20240716_212229627.jpg": SFW,
    "PXL_20240724_153838200.jpg": SFW,
    "PXL_20240724_154243335.jpg": SFW,
    "389497C6-8283-4F8C-A499-358845E9DE9E.jpg": SFW,
    "20260223_160230.jpg": SFW,
    "20260225_131350.jpg": SFW,
    "PXL_20240725_141202992.MP.jpg": SFW,
    "IMG_20250626_100625.jpg": SFW,
    "20250620_150027.jpg": SFW,
})

# Photo-and-tag posts (v2 look from photo_first.py). Key: (post id, slide number).
PICS = {
    ("x20261117", 3): lambda: pic("PXL_20240716_212229627.jpg", "Fair share.", T, BR, focus=(0.4, 0.55), size=90),
    ("x20261118", 1): lambda: pic("20250617_135603.jpg", "Compost Intensive Course", D, focus=(0.55, 0.5), logo=True),
    ("x20261118", 4): lambda: pic("IMG_4534.jpg", "Starts January 2027.", T, BR, focus=(0.5, 0.35)),
    ("ev-02", 3): lambda: pic("20250618_111904.jpg", "beopen.earth", G, focus=(0.45, 0.5)),
    ("ar20261123", 1): lambda: pic("20250625_155834.jpg", "Our first annual report", BR, focus=(0.3, 0.5), logo=True),
    ("ar20261123", 2): lambda: pic("20250626_093348.jpg", None, focus=(0.5, 0.5)),
    ("ar20261123", 3): lambda: pic("20250619_071702.jpg", "Coming soon.", T, BR, focus=(0.35, 0.5), size=90),
    ("x20261128", 1): lambda: pic("20250620_153134.jpg", "Small Business Saturday", T, BR, focus=(0.6, 0.45), logo=True),
    ("x20261128", 3): lambda: pic("20250620_144146.jpg", "Shop local for your soil.", D, focus=(0.35, 0.5)),
    ("x20261129", 2): lambda: pic("PXL_20240724_153838200.jpg", None, focus=(0.5, 0.5)),
    ("x20261129", 3): lambda: pic("389497C6-8283-4F8C-A499-358845E9DE9E.jpg", None, focus=(0.5, 0.5)),
    ("x20261130", 1): lambda: pic("R5A_3994.jpg", "10 weeks of compost.", G, focus=(0.5, 0.5), logo=True),
    ("x20261201", 2): lambda: pic("PXL_20240724_154243335.jpg", None, focus=(0.55, 0.5)),
    ("x20261201", 3): lambda: pic("20260223_160230.jpg", "Give a scholarship.", BR, focus=(0.4, 0.4)),
    ("x20261202", 1): lambda: pic("IMG_2593.jpg", "From shovel to slide.", BL, focus=(0.45, 0.6), logo=True),
    ("x20261202", 2): lambda: pic("20250311_112306.jpg", None, focus=(0.5, 0.35)),
    ("x20261202", 3): lambda: pic("20260225_131350.jpg", None, focus=(0.47, 0.4)),
    # Rebuild: soil aggregates, 4 slides.
    ("pl20261123", 1): lambda: pic("R5A_4016.jpg", "What holds soil together?", G, focus=(0.55, 0.5), logo=True, size=80),
    ("pl20261123", 2): lambda: pic("PXL_20240725_141202992.MP.jpg", "Roots and fungi.", T, BR, focus=(0.5, 0.45), size=90),
    ("pl20261123", 3): lambda: pic("IMG_20250626_100625.jpg", None, focus=(0.55, 0.5)),
    ("pl20261123", 4): lambda: pic("20250620_150027.jpg", "Tillage breaks it apart.", BR, focus=(0.5, 0.4)),
}

# Deck slides: (post id, slide n) -> (deck, deck slide, photo, focus, frame placement)
# Stories frames sit at (160, 0), 760 x 1350, on the light green ground; carousels fill 1080 x 1350.
STORIES = "26-10-2026-ig-stories-costa-rica-2027.pptx"
DECK = {
    ("st20261104", 1): (STORIES, 12, "IMG_4768.jpg", (0.5, 0.45), "story"),
    ("st20261116", 1): (STORIES, 5, "IMG_0087.jpg", (0.6, 0.45), "story"),
    ("st20261201", 1): (STORIES, 10, "20250314_094107.jpg", (0.22, 0.5), "story"),
    ("st20261201", 2): (STORIES, 11, "20250312_083929.jpg", (0.5, 0.3), "story"),
    ("cr20261112", 4): ("12-11-2026-costa-rica-a-day.pptx", 4, "20250310_163702.jpg", (0.5, 0.35), "full"),
}


def masks(deck, slides, work):
    """Render the deck with Picture 1 swapped for solid magenta: the magenta pixels are where
    the photo shows, so text, flowers and logo stay on top exactly as in the deck."""
    from pptx import Presentation
    mag = os.path.join(work, "magenta.png")
    Image.new("RGB", (64, 64), (255, 0, 255)).save(mag)
    p = Presentation(os.path.join(DECKS, deck))
    boxes = {}
    for n in slides:
        s = p.slides[n - 1]
        pic_ = [sh for sh in s.shapes if sh.name == "Picture 1"][0]
        _, rid = s.part.get_or_add_image_part(mag)
        pic_._element.xpath(".//a:blip")[0].set(
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed", rid)
        boxes[n] = (pic_.left, pic_.top, pic_.width, pic_.height, p.slide_width, p.slide_height)
    stem = os.path.splitext(deck)[0] + "-mask"
    path = os.path.join(work, stem + ".pptx")
    p.save(path)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", work, path],
                   check=True, capture_output=True, timeout=300)
    out = {}
    for n in slides:
        subprocess.run(["pdftoppm", "-r", "150", "-f", str(n), "-l", str(n), "-png",
                        os.path.join(work, stem + ".pdf"), os.path.join(work, f"{stem}-{n}")], check=True)
        f = [x for x in os.listdir(work) if x.startswith(f"{stem}-{n}-") and x.endswith(".png")][0]
        out[n] = (Image.open(os.path.join(work, f)).convert("RGB"), boxes[n])
    return out


def deck_slide(current, render, photo, focus, kind):
    tile = Image.open(os.path.join(TILES, current)).convert("RGB")
    fx, fy, fw, fh = (160, 0, 760, 1350) if kind == "story" else (0, 0, 1080, 1350)
    rend, (l, t, w, h, sw, sh) = render
    rend = rend.resize((fw, fh), Image.LANCZOS)
    k = fw / sw
    bx, by, bw, bh = round(l * k), round(t * k), round(w * k), round(h * k)
    ph = Image.open(os.path.join(B.SRC, photo)).convert("RGB")
    ph = ImageOps.fit(ph, (bw, bh), Image.LANCZOS, centering=focus)
    layer = Image.new("RGB", (fw, fh))
    layer.paste(ph, (bx, by))
    # magenta-ness -> mask (soft at anti-aliased edges)
    mask = Image.new("L", (fw, fh))
    px, mp = rend.load(), mask.load()
    for y in range(fh):
        for x in range(fw):
            R, G_, Bb = px[x, y]
            d = (255 - R) + G_ + (255 - Bb)
            mp[x, y] = 255 if d < 40 else (0 if d > 160 else int(255 * (160 - d) / 120))
    frame = tile.crop((fx, fy, fx + fw, fy + fh))
    frame = Image.composite(layer, frame, mask)
    tile.paste(frame, (fx, fy))
    return tile


def outname(pid, n):
    v = 1
    while os.path.exists(os.path.join(OUT, f"{pid}-c{v}-{n}.jpg")):
        v += 1
    return f"{pid}-c{v}-{n}.jpg"


def main():
    import json
    done = {}
    for (pid, n), make in PICS.items():
        name = outname(pid, n)
        make().save(os.path.join(OUT, name), quality=88, optimize=True)
        done[f"{pid}:{n}"] = name
    work = tempfile.mkdtemp(prefix="catalog_a_")
    by_deck = {}
    for key, (deck, sn, *_rest) in DECK.items():
        by_deck.setdefault(deck, []).append(sn)
    rendered = {deck: masks(deck, sorted(set(s)), work) for deck, s in by_deck.items()}
    current = {("st20261104", 1): "st20261104-v1-1.jpg", ("st20261116", 1): "st20261116-v1-1.jpg",
               ("st20261201", 1): "st20261201-v1-1.jpg", ("st20261201", 2): "st20261201-v1-2.jpg",
               ("cr20261112", 4): "cr20261112-v1-4.jpg"}
    for (pid, n), (deck, sn, photo, focus, kind) in DECK.items():
        img = deck_slide(current[(pid, n)], rendered[deck][sn], photo, focus, kind)
        name = outname(pid, n)
        img.save(os.path.join(OUT, name), quality=88, optimize=True)
        done[f"{pid}:{n}"] = name
    json.dump(done, open(os.environ.get("CA_JSON", "/dev/null"), "w"), indent=1)
    for k, v in done.items():
        print(k, v)


if __name__ == "__main__":
    main()
