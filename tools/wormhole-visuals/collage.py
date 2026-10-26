"""Layered collage kit for SFW posts, in the method of docs/visual-method/.

Real cutouts on paper backgrounds, one soft shadow, EB Garamond type, small
logo in a corner. Every deck renders to 1080 x 1350 JPGs and to an editable
.pptx where each background, cutout and text box is its own layer.

A deck is a strip of n slides. Slide k occupies x from k*W to (k+1)*W, so a
scene can run across slides (pass a layer wider than one slide) or each slide
can be its own picture.
"""
import io
import math
import os
import random
import re

import numpy as np
from PIL import Image, ImageDraw, ImageFilter, ImageFont

W, H = 1080, 1350
DEEP, GREEN, CREAM, KRAFT, SOIL, INK = "#22371F", "#156826", "#F4F1EA", "#CDB38A", "#4F3433", "#1E2A1B"
PURPLE = "#6B4C7A"  # Dr. Elaine memorial content only

LIB = os.environ.get("SFW_LIB", "/home/user/sfwlib")
FONTS = os.environ.get("SFW_FONTS", "/home/user/fonts")
FLOWERS = os.path.join(LIB, "social", "cutouts-flowers")
SOILC = os.path.join(LIB, "social", "cutouts-soil")
CUTS = os.environ.get("SFW_CUTS", "/home/user/sfwlib/cuts")  # rembg cutouts made for these posts
os.makedirs(CUTS, exist_ok=True)

FONT_FILES = {
    ("garamond", True): "eb-garamond-latin-500-italic.ttf",
    ("garamond", False): "eb-garamond-latin-400-normal.ttf",
    ("garamond-bold", False): "eb-garamond-latin-600-normal.ttf",
    ("montserrat", False): "Montserrat-SemiBold.ttf",
    ("montserrat-bold", False): "Montserrat-Bold.ttf",
    ("source", False): "SourceSans3-SemiBold.ttf",
    ("fredoka", False): "Fredoka-SemiBold.ttf",
    ("mont", False): "Montserrat-Regular.ttf",
    ("mont-medium", False): "Montserrat-SemiBold.ttf",
    ("pathway", False): "PathwayGothicOne-Regular.ttf",
    ("source-reg", False): "SourceSans3-Regular.ttf",
}
PPT_FONT = {"garamond": "EB Garamond", "garamond-bold": "EB Garamond", "montserrat": "Montserrat",
            "montserrat-bold": "Montserrat", "source": "Source Sans 3", "fredoka": "Fredoka",
            "mont": "Montserrat", "mont-medium": "Montserrat",
            "pathway": "Pathway Gothic One", "source-reg": "Source Sans 3"}
_EMOJI = re.compile("[\U0001F000-\U0001FAFF☀-➿️‍]")


def hexrgb(h):
    return tuple(int(h[i:i + 2], 16) for i in (1, 3, 5))


def font(kind, italic, size):
    return ImageFont.truetype(os.path.join(FONTS, FONT_FILES[(kind, italic)]), size)


# ---------------------------------------------------------------- pictures
def load(src):
    return src if isinstance(src, Image.Image) else Image.open(src).convert("RGB")


def cover(im, w, h, fx=0.5, fy=0.5):
    im = load(im)
    s = max(w / im.width, h / im.height)
    im = im.resize((max(w, round(im.width * s)), max(h, round(im.height * s))), Image.LANCZOS)
    x, y = round((im.width - w) * fx), round((im.height - h) * fy)
    return im.crop((x, y, x + w, y + h))


def paper(color, seed, w=W, h=H, grain=16):
    rng = np.random.default_rng(seed)
    base = np.array(hexrgb(color), np.float32)[None, None, :]
    n = rng.normal(0, 1, (max(2, h // 4), max(2, w // 4))).astype(np.float32)
    n = np.array(Image.fromarray(((n - n.min()) / np.ptp(n) * 255).astype(np.uint8)).resize((w, h), Image.BICUBIC), np.float32) / 255 - 0.5
    fine = rng.normal(0, 1, (h, w)).astype(np.float32)
    img = base + n[..., None] * grain + fine[..., None] * 3.5
    return Image.fromarray(np.clip(img, 0, 255).astype(np.uint8)).convert("RGBA")


def gradient(w, h, stops, seed=0):
    """Vertical painted gradient: stops is [(t, '#hex'), ...]; paper grain on top."""
    ts = np.array([t for t, _ in stops], np.float32)
    cs = np.array([hexrgb(c) for _, c in stops], np.float32)
    y = np.linspace(0, 1, h, dtype=np.float32)
    col = np.stack([np.interp(y, ts, cs[:, i]) for i in range(3)], 1)
    img = np.repeat(col[:, None, :], w, 1)
    rng = np.random.default_rng(seed)
    img += rng.normal(0, 2.5, (h, w, 1)).astype(np.float32)
    return Image.fromarray(np.clip(img, 0, 255).astype(np.uint8)).convert("RGBA")


def shadow(im, off=(10, 12), blur=11, alpha=0.4):
    a = im.getchannel("A")
    pad = blur * 3 + max(abs(off[0]), abs(off[1]))
    out = Image.new("RGBA", (im.width + 2 * pad, im.height + 2 * pad), (0, 0, 0, 0))
    sh = Image.new("RGBA", im.size, (30, 22, 15, 255))
    sh.putalpha(a.point(lambda v: int(v * alpha)))
    s = Image.new("RGBA", out.size, (0, 0, 0, 0))
    s.alpha_composite(sh, (pad + off[0], pad + off[1]))
    out.alpha_composite(s.filter(ImageFilter.GaussianBlur(blur)))
    out.alpha_composite(im, (pad, pad))
    return out, pad


def trim(im):
    box = im.getchannel("A").point(lambda v: 255 if v > 8 else 0).getbbox()
    return im.crop(box) if box else im


def defringe(im):
    """Pull the alpha edge in a little and drop faint pixels, so cutouts have no halo."""
    a = im.getchannel("A").filter(ImageFilter.MinFilter(3)).point(lambda v: 0 if v < 24 else v)
    im = im.copy()
    im.putalpha(a.filter(ImageFilter.GaussianBlur(0.6)))
    return im


def cutout(path, w=None, h=None, flip=False):
    im = trim(Image.open(path).convert("RGBA"))
    if flip:
        im = im.transpose(Image.FLIP_LEFT_RIGHT)
    if w:
        im = im.resize((w, max(1, round(im.height * w / im.width))), Image.LANCZOS)
    elif h:
        im = im.resize((max(1, round(im.width * h / im.height)), h), Image.LANCZOS)
    return im


def flower(name, **kw):
    return cutout(os.path.join(FLOWERS, name + ".png"), **kw)


def soil(name, **kw):
    return cutout(os.path.join(SOILC, name + ".png"), **kw)


_SESSION = None


def rembg_cut(src, name, crop=None, model="isnet-general-use"):
    """Cut a subject out of a real photo with rembg (cached as CUTS/name.png)."""
    out = os.path.join(CUTS, name + ".png")
    if not os.path.exists(out):
        global _SESSION
        from rembg import new_session, remove
        if _SESSION is None:
            _SESSION = new_session(model)
        im = Image.open(src).convert("RGB")
        if crop:
            im = im.crop(crop)
        if max(im.size) > 1800:
            im.thumbnail((1800, 1800), Image.LANCZOS)
        res = remove(im, session=_SESSION, post_process_mask=True)
        defringe(trim(res)).save(out)
    return out


def print_photo(src, w, h, fx=0.5, fy=0.5, border=16, color="#FBF8F1"):
    im = cover(src, w, h, fx, fy)
    out = Image.new("RGBA", (w + 2 * border, h + 2 * border), hexrgb(color) + (255,))
    out.paste(im, (border, border))
    return out


def round_photo(src, d, ring=12, fx=0.5, fy=0.5, ring_color=CREAM):
    im = cover(src, d, d, fx, fy).convert("RGBA")
    m = Image.new("L", (d * 2, d * 2), 0)
    ImageDraw.Draw(m).ellipse((0, 0, d * 2 - 1, d * 2 - 1), fill=255)
    im.putalpha(m.resize((d, d), Image.LANCZOS))
    out = Image.new("RGBA", (d + 2 * ring, d + 2 * ring), (0, 0, 0, 0))
    r = Image.new("L", ((d + 2 * ring) * 2,) * 2, 0)
    ImageDraw.Draw(r).ellipse((0, 0, r.width - 1, r.height - 1), fill=255)
    ring_im = Image.new("RGBA", out.size, hexrgb(ring_color) + (255,))
    ring_im.putalpha(r.resize(out.size, Image.LANCZOS))
    out.alpha_composite(ring_im)
    out.alpha_composite(im, (ring, ring))
    return out


def lens(src, d, fx=0.5, fy=0.5):
    """A micrograph seen through the eyepiece: round field, dark brass-black rim."""
    return round_photo(src, d, ring=max(10, d // 28), fx=fx, fy=fy, ring_color="#1B1B18")


def logo_disc(d=120):
    logo = Image.open(os.path.join(LIB, "img", "sfwlogo-240.png")).convert("RGBA")
    out = Image.new("RGBA", (d, d), (0, 0, 0, 0))
    m = Image.new("L", (d * 2, d * 2), 0)
    ImageDraw.Draw(m).ellipse((0, 0, d * 2 - 1, d * 2 - 1), fill=255)
    disc = Image.new("RGBA", (d, d), hexrgb(CREAM) + (255,))
    disc.putalpha(m.resize((d, d), Image.LANCZOS))
    out.alpha_composite(disc)
    l = logo.resize((round(d * 0.78), round(d * 0.78 * logo.height / logo.width)), Image.LANCZOS)
    out.alpha_composite(l, ((d - l.width) // 2, (d - l.height) // 2))
    return out


def phone(screenshot, w=430):
    """A plain dark phone body around a real page screenshot (top of the page)."""
    shot = load(screenshot)
    sw = w - 36
    sh = round(sw * 2.05)
    shot = shot.resize((sw, round(shot.height * sw / shot.width)), Image.LANCZOS).crop((0, 0, sw, sh))
    body = Image.new("RGBA", (w, sh + 60), (0, 0, 0, 0))
    m = Image.new("L", (body.width * 2, body.height * 2), 0)
    ImageDraw.Draw(m).rounded_rectangle((0, 0, m.width - 1, m.height - 1), radius=120, fill=255)
    dark = Image.new("RGBA", body.size, (24, 26, 22, 255))
    dark.putalpha(m.resize(body.size, Image.LANCZOS))
    body.alpha_composite(dark)
    sm = Image.new("L", (sw * 2, sh * 2), 0)
    ImageDraw.Draw(sm).rounded_rectangle((0, 0, sm.width - 1, sm.height - 1), radius=70, fill=255)
    s = shot.convert("RGBA")
    s.putalpha(sm.resize((sw, sh), Image.LANCZOS))
    body.alpha_composite(s, (18, 30))
    return body


# ---------------------------------------------------------------- drawn things
def draw_layer(w, h, fn, scale=2):
    """Draw at 2x with fn(ImageDraw, scale) and shrink, for smooth hand-drawn elements."""
    im = Image.new("RGBA", (w * scale, h * scale), (0, 0, 0, 0))
    fn(ImageDraw.Draw(im), scale)
    return im.resize((w, h), Image.LANCZOS)


def hyphae(w, h, seed, color=CREAM, n=6, width=2.2, start=None):
    """Fine branching fungal threads, drawn."""
    rnd = random.Random(seed)

    def fn(d, s):
        def grow(x, y, ang, length, wd, depth):
            pts = [(x, y)]
            for _ in range(int(length / 6)):
                ang += rnd.uniform(-0.22, 0.22)
                x += 6 * math.cos(ang)
                y += 6 * math.sin(ang)
                pts.append((x, y))
                if depth < 3 and rnd.random() < 0.05:
                    grow(x, y, ang + rnd.choice((-1, 1)) * rnd.uniform(0.5, 1.1), length * 0.45, wd * 0.7, depth + 1)
            d.line([(px * s, py * s) for px, py in pts], fill=hexrgb(color) + (230,), width=max(1, round(wd * s)), joint="curve")

        for _ in range(n):
            x, y = start if start else (rnd.uniform(0, w), rnd.uniform(0, h))
            grow(x, y, rnd.uniform(0, 6.28), rnd.uniform(0.3, 0.7) * max(w, h), width, 0)

    return draw_layer(w, h, fn)


def roots(w, h, seed, color="#E9DCC3", n=7, width=3.0):
    """Drawn roots hanging down from the top edge of the layer."""
    rnd = random.Random(seed)

    def fn(d, s):
        def grow(x, y, ang, length, wd, depth):
            pts = [(x, y)]
            for _ in range(int(length / 7)):
                ang += rnd.uniform(-0.18, 0.18)
                ang = max(0.4, min(2.7, ang))
                x += 7 * math.cos(ang)
                y += 7 * math.sin(ang)
                pts.append((x, y))
                if depth < 3 and rnd.random() < 0.07:
                    grow(x, y, ang + rnd.choice((-1, 1)) * rnd.uniform(0.4, 0.9), length * 0.5, wd * 0.65, depth + 1)
            d.line([(px * s, py * s) for px, py in pts], fill=hexrgb(color) + (235,), width=max(1, round(wd * s)), joint="curve")

        for i in range(n):
            grow(w * (i + 0.5) / n + rnd.uniform(-30, 30), 0, math.pi / 2 + rnd.uniform(-0.3, 0.3), rnd.uniform(0.5, 0.95) * h, width, 0)

    return draw_layer(w, h, fn)


def rain(w, h, seed, color="#DDE6E2", n=140):
    rnd = random.Random(seed)

    def fn(d, s):
        for _ in range(n):
            x, y, ln = rnd.uniform(0, w), rnd.uniform(0, h), rnd.uniform(24, 60)
            d.line((x * s, y * s, (x - ln * 0.18) * s, (y + ln) * s), fill=hexrgb(color) + (rnd.randint(90, 170),), width=max(1, round(1.6 * s)))

    return draw_layer(w, h, fn)


def card(w, h, color="#FBF8F1"):
    out = Image.new("RGBA", (w, h), hexrgb(color) + (255,))
    return out


# ---------------------------------------------------------------- deck model
class Deck:
    def __init__(self, post_id, n, slug, iso):
        self.id, self.n, self.slug, self.iso = post_id, n, slug, iso
        self.items = []  # dicts in strip coordinates
        self.notes = []

    # layers ------------------------------------------------------
    def bg(self, k, im, name="background"):
        """Background for slide k (or a strip-wide one when k is None)."""
        x = 0 if k is None else k * W
        self.items.append(dict(type="image", name=name, im=im.convert("RGBA"), x=x, y=0, bg=True))

    def put(self, k, name, im, cx, cy, rot=0, shadowed=True, anchor="center", off=(10, 12), blur=11, alpha=0.4):
        """Place an RGBA picture on slide k; cx, cy are slide coordinates of its centre."""
        if rot:
            im = im.rotate(rot, expand=True, resample=Image.BICUBIC)
        pad = 0
        if shadowed:
            im, pad = shadow(im, off, blur, alpha)
        if anchor == "tl":
            x, y = cx - pad, cy - pad
        else:
            x, y = round(cx - im.width / 2), round(cy - im.height / 2)
        self.items.append(dict(type="image", name=name, im=im, x=x + k * W, y=y, bg=False))

    def text(self, k, name, text, x, y, w, h, size, kind="garamond", italic=True, color=DEEP, lh=1.08,
             align="left", min_size=None, shade=False, track=0):
        text = _EMOJI.sub("", text).strip()
        f, lines, size = fit(text, kind, italic, w, h, size, min_size or max(18, int(size * 0.55)), lh, track)
        self.items.append(dict(type="text", name=name, text=text, lines=lines, x=x + k * W, y=y, w=w, h=h, size=size,
                               kind=kind, italic=italic, color=color, lh=lh, align=align, shade=shade, track=track))
        return y + len(lines) * size * lh  # bottom of the block

    def logo(self, k, corner="tr", d=112):
        pos = {"tr": (W - 92, 92), "tl": (92, 92), "br": (W - 92, H - 92), "bl": (92, H - 92)}[corner]
        self.put(k, "logo", logo_disc(d), *pos, off=(4, 6), blur=6, alpha=0.3)

    # render ------------------------------------------------------
    def render(self):
        strip = Image.new("RGBA", (W * self.n, H), (255, 255, 255, 255))
        for it in self.items:
            if it["type"] == "image":
                strip.alpha_composite(_clip(it["im"], it["x"], it["y"], strip.size), (max(0, it["x"]), max(0, it["y"])))
        for it in self.items:
            if it["type"] == "text":
                draw_text(strip, it)
        return [strip.crop((k * W, 0, (k + 1) * W, H)).convert("RGB") for k in range(self.n)]

    def slide_texts(self):
        out = []
        for k in range(self.n):
            t = [it["text"] for it in self.items if it["type"] == "text" and k * W <= it["x"] < (k + 1) * W]
            out.append(" ".join(t))
        return out

    def save(self, out_dir, pptx_dir=None):
        os.makedirs(out_dir, exist_ok=True)
        slides = self.render()
        names = []
        for k, im in enumerate(slides):
            # a "-v2" slug marks a redesign: new file names, so live files are never overwritten
            ver = "-" + self.slug.rsplit("-", 1)[1] if re.search(r"-v\d+$", self.slug) else ""
            fn = f"{self.id}{ver}-{k + 1}.jpg"
            im.save(os.path.join(out_dir, fn), quality=84, optimize=True, progressive=True)
            names.append(fn)
        if pptx_dir:
            self.pptx(os.path.join(pptx_dir, self.pptx_name()))
        return names, slides

    def pptx_name(self):
        y, m, d = self.iso.split("-")
        return f"{d}-{m}-{y}-{self.slug}.pptx"

    def pptx(self, path):
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Emu, Pt

        EMU = 914400 / 108
        prs = Presentation()
        prs.slide_width, prs.slide_height = Emu(round(W * EMU)), Emu(round(H * EMU))
        for k in range(self.n):
            sl = prs.slides.add_slide(prs.slide_layouts[6])
            x0 = k * W
            for it in self.items:
                if it["type"] != "image":
                    continue
                x, y, im = it["x"] - x0, it["y"], it["im"]
                if x >= W or x + im.width <= 0:
                    continue
                box = (max(0, -x), max(0, -y), min(im.width, W - x), min(im.height, H - y))
                piece = im.crop(box)
                if piece.getchannel("A").getbbox() is None:
                    continue
                buf = io.BytesIO()
                if it["bg"] or piece.getchannel("A").getextrema()[0] == 255:
                    piece.convert("RGB").save(buf, "JPEG", quality=86)
                else:
                    piece.save(buf, "PNG", optimize=True)
                buf.seek(0)
                pic = sl.shapes.add_picture(buf, Emu(round((x + box[0]) * EMU)), Emu(round((y + box[1]) * EMU)),
                                            Emu(round(piece.width * EMU)), Emu(round(piece.height * EMU)))
                pic.name = it["name"]
            for it in self.items:
                if it["type"] != "text" or not (x0 <= it["x"] < x0 + W):
                    continue
                tb = sl.shapes.add_textbox(Emu(round((it["x"] - x0) * EMU)), Emu(round(it["y"] * EMU)),
                                           Emu(round(it["w"] * EMU)), Emu(round(it["h"] * EMU)))
                tb.name = it["name"]
                tf = tb.text_frame
                tf.word_wrap = True
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                p = tf.paragraphs[0]
                p.line_spacing = it["lh"]
                p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[it["align"]]
                r = p.add_run()
                r.text = it["text"]
                f = r.font
                f.name = PPT_FONT[it["kind"]]
                f.size = Pt(round(it["size"] * 72 / 108, 1))
                f.italic = it["italic"]
                f.bold = it["kind"] in ("garamond-bold", "montserrat", "montserrat-bold", "source", "fredoka", "mont-medium")
                f.color.rgb = RGBColor(*hexrgb(it["color"]))
        prs.save(path)


def _clip(im, x, y, size):
    box = (max(0, -x), max(0, -y), min(im.width, size[0] - x), min(im.height, size[1] - y))
    return im.crop(box)


# ---------------------------------------------------------------- type
def _wrap(text, f, w, track):
    lines = []
    for para in text.split("\n"):
        cur = ""
        for word in para.split(" "):
            t = (cur + " " + word).strip()
            if _width(t, f, track) <= w or not cur:
                cur = t
            else:
                lines.append(cur)
                cur = word
        lines.append(cur)
    return lines


def _width(t, f, track):
    return f.getlength(t) + track * max(0, len(t) - 1)


def fit(text, kind, italic, w, h, size, min_size, lh, track=0):
    s = size
    while True:
        f = font(kind, italic, s)
        lines = _wrap(text, f, w, track)
        if (len(lines) * s * lh <= h and all(_width(l, f, track) <= w for l in lines)) or s <= min_size:
            if s <= min_size and len(lines) * s * lh > h:
                raise ValueError(f"Text does not fit its box: {text[:60]!r}")
            return f, lines, s
        s -= 2


def draw_text(strip, it):
    f = font(it["kind"], it["italic"], it["size"])
    layer = Image.new("RGBA", strip.size, (0, 0, 0, 0))
    d = ImageDraw.Draw(layer)
    y = it["y"]
    step = it["size"] * it["lh"]
    for line in it["lines"]:
        lw = _width(line, f, it["track"])
        x = it["x"] + {"left": 0, "center": (it["w"] - lw) / 2, "right": it["w"] - lw}[it["align"]]
        if it["track"]:
            cx = x
            for ch in line:
                d.text((cx, y), ch, font=f, fill=hexrgb(it["color"]) + (255,))
                cx += f.getlength(ch) + it["track"]
        else:
            d.text((x, y), line, font=f, fill=hexrgb(it["color"]) + (255,))
        y += step
    if it["shade"]:
        a = layer.getchannel("A")
        sh = Image.new("RGBA", strip.size, (10, 16, 10, 0))
        sh.putalpha(a.point(lambda v: int(v * 0.55)))
        sh = sh.filter(ImageFilter.GaussianBlur(5))
        strip.alpha_composite(sh, (0, 0)) if False else strip.alpha_composite(_shift(sh, 2, 3))
    strip.alpha_composite(layer)


def _shift(im, dx, dy):
    out = Image.new("RGBA", im.size, (0, 0, 0, 0))
    out.paste(im, (dx, dy))
    return out
