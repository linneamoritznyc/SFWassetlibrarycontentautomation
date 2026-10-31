"""Every news post in Wormhole in one PowerPoint for Allison (Linnea, 31 Oct): a divider slide per story,
then its slides. Stories with an editable deck keep their editable layers; the rest go in as their tiles.

  python3 tools/wormhole-visuals/news_deck.py <posts.json> <out.pptx>
"""
import copy
import io
import json
import os
import re
import sys
from datetime import date

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
PUB = os.path.abspath(os.path.join(HERE, "..", "..", "apps", "web", "public"))
DECKS = os.path.join(PUB, "wormhole", "decks")
PX = 9525
W, H = 1080, 1350
GREEN, CREAM = RGBColor(0x15, 0x68, 0x26), RGBColor(0xF4, 0xF1, 0xEA)
EDITABLE = {  # post id -> its editable deck
    "news-1000farms": "20-11-2026-regenerative-farms-39-percent-carbon-v5.pptx",
    "pl20261127": "27-11-2026-fun-soil-news-november-v1.pptx",
    "pl20261218": "18-12-2026-fun-soil-news-december-v1.pptx",
}
R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"


def is_news(p):
    return bool(re.search(r"\bnews\b|shoutout|new study|research pick", p["topic"], re.I)
                or re.search(r"news card|article repost card", p["visual_note"], re.I))


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    for shp in list(s.shapes):
        shp._element.getparent().remove(shp._element)
    return s


def textbox(s, name, text, y, h, size, color, bold=False):
    tb = s.shapes.add_textbox(Emu(90 * PX), Emu(y * PX), Emu(900 * PX), Emu(h * PX))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    r = tf.paragraphs[0].add_run()
    r.text = text
    r.font.name = "Montserrat" if bold else "Source Sans 3"
    r.font.size, r.font.bold = Pt(size * 0.75), bold
    r.font.color.rgb = color


def divider(prs, p, n, kind):
    s = blank(prs)
    bg = s.shapes.add_shape(1, 0, 0, Emu(W * PX), Emu(H * PX))
    bg.name = "Background"
    bg.fill.solid()
    bg.fill.fore_color.rgb = GREEN
    bg.line.fill.background()
    day = "Date TBD" if p["date_tbd"] else date.fromisoformat(p["iso_date"]).strftime("%a %-d %b %Y")
    textbox(s, "Date", f"{day} · {p['platform']}", 420, 120, 52, CREAM)
    textbox(s, "Topic", p["topic"], 560, 420, 72, CREAM, bold=True)
    textbox(s, "Slides", f"{n} slide{'s' if n != 1 else ''} follow · {kind} · Wormhole post {p['id']}", 1050, 160,
            42, CREAM)
    s.notes_slide.notes_text_frame.text = p["caption"]


def image_slide(prs, path, k, p):
    s = blank(prs)
    im = Image.open(path)
    sc = min(W / im.width, H / im.height)
    w, h = round(im.width * sc), round(im.height * sc)
    if (w, h) != (W, H):  # a LinkedIn 1200 x 627 card sits centred on cream
        bg = s.shapes.add_shape(1, 0, 0, Emu(W * PX), Emu(H * PX))
        bg.name = "Background"
        bg.fill.solid()
        bg.fill.fore_color.rgb = CREAM
        bg.line.fill.background()
    pic = s.shapes.add_picture(path, Emu((W - w) // 2 * PX), Emu((H - h) // 2 * PX), Emu(w * PX), Emu(h * PX))
    pic.name = f"Slide {k} image"
    s.notes_slide.notes_text_frame.text = p["slides"][k - 1]["text"] if k <= len(p["slides"]) else ""


def copy_deck(prs, path):
    """Each slide of an editable deck, layer by layer: pictures re-added from their bytes, the rest copied."""
    src = Presentation(path)
    sx, sy = prs.slide_width / src.slide_width, prs.slide_height / src.slide_height
    for ss in src.slides:
        s = blank(prs)
        for shp in ss.shapes:
            if shp.shape_type == 13:  # picture
                pic = s.shapes.add_picture(io.BytesIO(shp.image.blob), round(shp.left * sx), round(shp.top * sy),
                                           round(shp.width * sx), round(shp.height * sy))
                pic.name = shp.name
            else:
                el = copy.deepcopy(shp._element)
                if any(n.get(R_EMBED) for n in el.iter()):
                    raise SystemExit(f"{path}: shape {shp.name} links an image; copy not supported")
                s.shapes._spTree.append(el)
        if ss.has_notes_slide:
            s.notes_slide.notes_text_frame.text = ss.notes_slide.notes_text_frame.text


def main(posts_json, out):
    d = json.load(open(posts_json))
    posts = d["posts"] if isinstance(d, dict) else d
    news = [p for p in posts if is_news(p) and any(s["image"] for s in p["slides"])]
    ids = {p["id"] for p in news}
    news = [p for p in news if not (p["id"].endswith("-li") and p["id"][:-3] in ids)]  # same card as the IG post
    news.sort(key=lambda p: (p["date_tbd"], p["iso_date"], p["id"]))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W * PX), Emu(H * PX)
    for p in news:
        if p["id"] in EDITABLE:
            src = os.path.join(DECKS, EDITABLE[p["id"]])
            divider(prs, p, len(Presentation(src).slides), "editable")
            copy_deck(prs, src)
        else:
            imgs = [s["image"] for s in p["slides"] if s["image"] and not s["text"].lower().startswith("cover")]
            divider(prs, p, len(imgs), "images")
            for k, im in enumerate(imgs, 1):
                image_slide(prs, os.path.join(PUB, im.lstrip("/")), k, p)
        print(p["id"], p["iso_date"], p["topic"][:50])
    prs.save(out)
    print(out, len(prs.slides), "slides")


if __name__ == "__main__":
    main(*sys.argv[1:3])
