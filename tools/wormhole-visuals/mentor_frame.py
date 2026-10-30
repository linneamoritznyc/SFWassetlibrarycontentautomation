"""Mentor profile images in Linnea's flower frame (her Canva deck Mentor-Profile-Images-NoText), with names.

Takes her one-slide deck (Casey in the flower frame), copies the slide once per mentor, swaps the framed
photo (a JPG with a thin cream border, filling the frame shape) for that mentor's photo, and adds the name
and role as real text boxes in the white space under the frame. Every cutout group stays as she made it.

  python3 tools/wormhole-visuals/mentor_frame.py <her deck.pptx> <photo dir> <out.pptx>
    photo dir holds photo1.png ... photo5.png in the order of MENTORS below.
"""
import copy
import io
import os
import sys

from PIL import Image, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

MENTORS = [  # (name, role) in the order of the photos
    ("Casey Williams", "Soil Food Web mentor"),
    ("Brian Daubenspeck", "Soil Food Web mentor"),
    ("Delvin Solkinson", "Permaculture Lead"),
    ("Tommy Tepper", "Director of Education"),
    ("Dora Tkalec", "Soil Food Web mentor"),
]
PX = 9525  # EMU per pixel; the deck is 1080 x 1350
BORDER, CREAM = 14, (251, 248, 241)
SOIL, GREEN = (0x4F, 0x34, 0x33), (0x15, 0x68, 0x26)
R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"


def framed_photo(src, w=734, h=747):
    """The mentor's photo cropped to the frame, inside the same thin cream border as the original."""
    im = ImageOps.exif_transpose(Image.open(src)).convert("RGB")
    iw, ih = w - 2 * BORDER, h - 2 * BORDER
    s = max(iw / im.width, ih / im.height)
    im = im.resize((round(im.width * s), round(im.height * s)), Image.LANCZOS)
    x, y = (im.width - iw) // 2, max(0, round((im.height - ih) * 0.35))
    out = Image.new("RGB", (w, h), CREAM)
    out.paste(im.crop((x, y, x + iw, y + ih)), (BORDER, BORDER))
    buf = io.BytesIO()
    out.save(buf, "JPEG", quality=92)
    buf.seek(0)
    return buf


def copy_slide(prs, src):
    """A new slide with every shape of src, pictures re-linked to the same image parts."""
    new = prs.slides.add_slide(prs.slide_layouts[6])
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)
    rid_map = {}
    for rid, rel in src.part.rels.items():
        if "image" in rel.reltype:
            rid_map[rid] = new.part.relate_to(rel.target_part, rel.reltype)
    for el in src.shapes._spTree.iterchildren():
        if el.tag.endswith("}nvGrpSpPr") or el.tag.endswith("}grpSpPr"):
            continue
        el2 = copy.deepcopy(el)
        for node in el2.iter():
            if node.get(R_EMBED) in rid_map:
                node.set(R_EMBED, rid_map[node.get(R_EMBED)])
        new.shapes._spTree.append(el2)
    return new


def photo_shape(slide):
    for shp in slide.shapes:
        if shp.shape_type == 6 and shp.name == "Group 54":
            return shp.shapes[0]
    raise SystemExit("framed photo (Group 54) not found")


def add_text(slide, text, y, h, size, color, bold, font):
    tb = slide.shapes.add_textbox(Emu(60 * PX), Emu(y * PX), Emu(960 * PX), Emu(h * PX))
    tb.name = text
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name, r.font.size, r.font.bold = font, Pt(size * 0.75), bold
    r.font.color.rgb = RGBColor(*color)


def main(deck, photos, out):
    prs = Presentation(deck)
    base = prs.slides[0]
    for i, (name, role) in enumerate(MENTORS, start=1):
        sl = copy_slide(prs, base)
        blip = photo_shape(sl)._element.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip")
        _, rid = sl.part.get_or_add_image_part(framed_photo(os.path.join(photos, f"photo{i}.png")))
        blip.set(R_EMBED, rid)
        photo_shape(sl)._element.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr").set(
            "descr", f"Photo: {name}")
        add_text(sl, name, 1222, 76, 62, SOIL, True, "Montserrat")
        add_text(sl, role, 1298, 44, 34, GREEN, False, "Source Sans 3")
        sl.notes_slide.notes_text_frame.text = f"{name}, {role}. Flower frame from Linnea's Canva deck."
    # drop the original no-text slide
    sld = prs.slides._sldIdLst[0]
    prs.part.drop_rel(sld.rId)
    prs.slides._sldIdLst.remove(sld)
    prs.save(out)


if __name__ == "__main__":
    main(*sys.argv[1:4])
