"""One continuous sunset landscape across six slides, with a soil cross-section below."""
import os, math, random
import numpy as np
from PIL import Image, ImageDraw, ImageFont, ImageFilter
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR

D = '/home/claude/p05/fl/'
M = '/home/claude/sm/assets/microscopy/'
OUT = '/home/claude/p05/v3/'
os.makedirs(OUT + 'layers', exist_ok=True)
N = 6
W, H = 1080, 1350
PW = W * N
GROUND = 820          # ground line
FONTS = os.path.expanduser('~/.fonts/')
FI = lambda s: ImageFont.truetype(FONTS + 'eb-garamond-latin-500-italic.ttf', s)
FR = lambda s: ImageFont.truetype(FONTS + 'eb-garamond-latin-400-normal.ttf', s)
CREAM = (250, 244, 230)
random.seed(7); np.random.seed(7)


def grad(h, stops):
    ys = np.linspace(0, 1, h)
    out = np.zeros((h, 3))
    for c in range(3):
        out[:, c] = np.interp(ys, [s[0] for s in stops], [s[1][c] for s in stops])
    return out


def ridge(base, amp, freq, seed, w=PW):
    rng = np.random.default_rng(seed)
    x = np.arange(w)
    y = np.full(w, float(base))
    for k in range(1, 5):
        ph = rng.uniform(0, 6.28)
        y += amp / k * np.sin(x / w * 2 * math.pi * freq * k + ph)
    return y


def fill_under(img, ys, color, alpha=255, blur=0):
    lay = Image.new('RGBA', img.size, (0, 0, 0, 0))
    d = ImageDraw.Draw(lay)
    pts = [(i, ys[i]) for i in range(0, len(ys), 4)] + [(len(ys) - 1, ys[-1]), (len(ys) - 1, img.height), (0, img.height)]
    d.polygon(pts, fill=color + (alpha,))
    if blur: lay = lay.filter(ImageFilter.GaussianBlur(blur))
    img.alpha_composite(lay)


def shadow(im, off=(10, 14), blur=12, alpha=0.42):
    pad = blur * 3 + max(off)
    out = Image.new('RGBA', (im.width + pad * 2, im.height + pad * 2), (0, 0, 0, 0))
    s = Image.new('RGBA', im.size, (20, 14, 8, 0)); s.putalpha(im.split()[3].point(lambda v: int(v * alpha)))
    sl = Image.new('RGBA', out.size, (0, 0, 0, 0)); sl.alpha_composite(s, (pad + off[0], pad + off[1]))
    out.alpha_composite(sl.filter(ImageFilter.GaussianBlur(blur))); out.alpha_composite(im, (pad, pad))
    return out, pad


def cut(path, w=None, h=None, rot=0, flip=False):
    im = Image.open(path).convert('RGBA')
    if flip: im = im.transpose(Image.FLIP_LEFT_RIGHT)
    if w: im = im.resize((w, max(1, int(im.height * w / im.width))), Image.LANCZOS)
    elif h: im = im.resize((max(1, int(im.width * h / im.height)), h), Image.LANCZOS)
    if rot: im = im.rotate(rot, expand=True, resample=Image.BICUBIC)
    return im


def place(img, im, x, y, sh=True, anchor='bottom'):
    """x = centre, y = bottom of the cutout."""
    if sh:
        big, pad = shadow(im)
    else:
        big, pad = im, 0
    tx = int(x - im.width / 2 - pad); ty = int(y - im.height - pad) if anchor == 'bottom' else int(y - pad)
    img.alpha_composite(big, (max(tx, 0), max(ty, 0))) if tx >= 0 and ty >= 0 else img.paste(big, (tx, ty), big)


# ---------- sky ----------
sky = grad(GROUND + 40, [(0, (58, 52, 96)), (0.35, (150, 96, 120)), (0.62, (236, 150, 104)), (0.85, (250, 200, 140)), (1, (252, 226, 180))])
P = Image.fromarray(np.repeat(sky[:, None, :], PW, axis=1).astype(np.uint8)).convert('RGBA')
P = P.resize((PW, GROUND + 40))
canvas = Image.new('RGBA', (PW, H), (0, 0, 0, 255)); canvas.alpha_composite(P, (0, 0))

# sun with soft glow, on slide 1-2 seam area
glow = Image.new('RGBA', (PW, H), (0, 0, 0, 0)); g = ImageDraw.Draw(glow)
sx, sy = 1500, 640
for r, a in [(520, 40), (380, 60), (260, 90), (170, 140)]:
    g.ellipse((sx - r, sy - r, sx + r, sy + r), fill=(255, 214, 150, a))
glow = glow.filter(ImageFilter.GaussianBlur(60)); canvas.alpha_composite(glow)
sun = Image.new('RGBA', (PW, H), (0, 0, 0, 0)); ImageDraw.Draw(sun).ellipse((sx - 120, sy - 120, sx + 120, sy + 120), fill=(255, 236, 196, 255))
canvas.alpha_composite(sun.filter(ImageFilter.GaussianBlur(3)))

# birds, drawn
d = ImageDraw.Draw(canvas)
for bx, by, s in [(1800, 380, 22), (1850, 410, 16), (1760, 430, 14), (4300, 300, 20), (4350, 330, 14)]:
    d.arc((bx - s, by - s / 2, bx, by + s / 2), 200, 340, fill=(60, 40, 60, 255), width=3)
    d.arc((bx, by - s / 2, bx + s, by + s / 2), 200, 340, fill=(60, 40, 60, 255), width=3)

# hills, far to near, with haze fades
fill_under(canvas, ridge(600, 50, 1.3, 1), (186, 128, 140), 200, blur=2)
fill_under(canvas, ridge(680, 45, 2.1, 2), (128, 104, 112), 230, blur=1)
haze = Image.new('RGBA', canvas.size, (0, 0, 0, 0)); hz = ImageDraw.Draw(haze)
for i in range(60):
    hz.line((0, 640 + i * 2, PW, 640 + i * 2), fill=(250, 210, 170, max(0, 70 - i)))
canvas.alpha_composite(haze)


# trees: drawn, recursive branches + foliage, silhouette darkened by the light
def tree(img, x, base, h, seed, tone=(48, 52, 40), leaf=(70, 86, 52)):
    r = random.Random(seed)
    lay = Image.new('RGBA', img.size, (0, 0, 0, 0)); t = ImageDraw.Draw(lay)

    def br(x0, y0, ang, ln, wd, depth):
        x1 = x0 + ln * math.cos(ang); y1 = y0 - ln * math.sin(ang)
        t.line((x0, y0, x1, y1), fill=tone + (255,), width=max(1, int(wd)))
        if depth == 0 or ln < 14:
            for _ in range(3):
                rr = r.uniform(18, 46) * h / 520
                ox, oy = x1 + r.uniform(-22, 22), y1 + r.uniform(-22, 22)
                t.ellipse((ox - rr, oy - rr, ox + rr, oy + rr), fill=leaf + (r.randint(150, 220),))
            return
        for da in (r.uniform(0.25, 0.55), -r.uniform(0.25, 0.55)) + ((r.uniform(-0.15, 0.15),) if r.random() < 0.5 else ()):
            br(x1, y1, ang + da, ln * r.uniform(0.68, 0.8), wd * 0.68, depth - 1)

    br(x, base, math.pi / 2 + r.uniform(-0.05, 0.05), h * 0.32, h / 26, 7)
    sh = lay.filter(ImageFilter.GaussianBlur(1))
    img.alpha_composite(sh)


for x, hgt, sd, tone in [(420, 420, 3, (60, 58, 52)), (2700, 360, 5, (70, 64, 60)), (3600, 520, 9, (44, 48, 36)), (5600, 460, 11, (52, 54, 42))]:
    tree(canvas, x, GROUND + 10, hgt, sd, tone=tone)

# near meadow band
fill_under(canvas, ridge(GROUND - 12, 10, 3.0, 4), (92, 112, 58), 255)
d = ImageDraw.Draw(canvas)
for i in range(2600):
    x = random.randint(0, PW); y = GROUND + random.randint(-14, 8)
    d.line((x, y, x + random.randint(-6, 6), y - random.randint(10, 34)), fill=(random.randint(60, 110), random.randint(100, 140), random.randint(40, 70), 255), width=2)

# fence, hand drawn, across slides 3 to 5
fx0, fx1 = 2350, 4950
for x in range(fx0, fx1 + 1, 170):
    j = random.randint(-6, 6)
    d.line((x + j, GROUND - 150, x, GROUND + 6), fill=(92, 70, 52, 255), width=10)
    d.line((x + j - 4, GROUND - 150, x + j + 4, GROUND - 150), fill=(70, 52, 40, 255), width=6)
for yy in (GROUND - 120, GROUND - 60):
    pts = [(x, yy + 4 * math.sin(x / 90)) for x in range(fx0, fx1 + 1, 20)]
    d.line(pts, fill=(110, 84, 60, 255), width=8)

# ---------- soil cross-section ----------
soil = grad(H - GROUND, [(0, (78, 52, 40)), (0.3, (62, 40, 32)), (1, (34, 22, 20))])
S = Image.fromarray(np.repeat(soil[:, None, :], PW, axis=1).astype(np.uint8))
noise = (np.random.rand(H - GROUND, PW) * 26 - 13)
S = Image.fromarray(np.clip(np.array(S, np.float32) + noise[:, :, None], 0, 255).astype(np.uint8)).convert('RGBA')
canvas.alpha_composite(S, (0, GROUND + 8))
d = ImageDraw.Draw(canvas)
d.line([(x, GROUND + 8 + 5 * math.sin(x / 70)) for x in range(0, PW, 10)], fill=(120, 90, 60, 255), width=6)
for i in range(900):  # pebbles and crumbs
    x = random.randint(0, PW); y = random.randint(GROUND + 30, H); r = random.randint(2, 8)
    c = random.choice([(110, 84, 66), (90, 70, 56), (140, 118, 96), (54, 38, 32)])
    d.ellipse((x - r, y - r, x + r, y + r * 0.7), fill=c + (255,))


# roots and fungal hyphae as fine drawn lines
def root(x, y, ang, ln, wd, depth, col, r):
    if depth == 0 or ln < 6: return
    x1 = x + ln * math.cos(ang); y1 = y + ln * math.sin(ang)
    d.line((x, y, x1, y1), fill=col, width=max(1, int(wd)))
    for da in (r.uniform(-0.6, -0.15), r.uniform(0.15, 0.6)):
        root(x1, y1, ang + da, ln * r.uniform(0.6, 0.8), wd * 0.7, depth - 1, col, r)


rr = random.Random(3)
for x in range(80, PW, 260):
    root(x + rr.randint(-40, 40), GROUND + 10, math.pi / 2 + rr.uniform(-0.3, 0.3), rr.randint(90, 150), 6, 7, (196, 168, 130, 255), rr)
hy = Image.new('RGBA', canvas.size, (0, 0, 0, 0)); hd = ImageDraw.Draw(hy)
for k in range(260):  # hyphae web, densest on slides 4 and 5
    cx = rr.choice([rr.randint(0, PW), rr.randint(3240, 5400), rr.randint(3240, 5400)])
    cy = rr.randint(GROUND + 60, H - 30)
    pts = [(cx, cy)]
    a = rr.uniform(0, 6.28)
    for _ in range(rr.randint(8, 20)):
        a += rr.uniform(-0.5, 0.5); cx += 18 * math.cos(a); cy += 18 * math.sin(a) * 0.6
        pts.append((cx, cy))
    hd.line(pts, fill=(246, 238, 220, rr.randint(90, 170)), width=2)
canvas.alpha_composite(hy.filter(ImageFilter.GaussianBlur(0.4)))

# microscope organisms in the soil, small
for i, (f, x, y, w) in enumerate([('cutout-2.png', 2350, 1150, 170), ('cutout-3.png', 2700, 1230, 160), ('cutout-6.png', 2950, 1120, 180),
                                   ('cutout-1.png', 3900, 1240, 110), ('cutout-7.png', 4500, 1180, 190), ('cutout-5.png', 5900, 1200, 200)]):
    o = cut(M + f, w=w, rot=random.randint(-20, 20))
    place(canvas, o, x, y, sh=False)

# ---------- objects on the ground ----------
place(canvas, cut(D + 's8-03-wood-chips.png', w=420), 1600, GROUND + 50)
place(canvas, cut(D + 's4-04-moss.png', w=320, flip=True), 1900, GROUND + 50)
place(canvas, cut(D + 's8-02-compost-fungal-threads.png', w=300), 2950, GROUND + 60)
place(canvas, cut(D + 's8-04-dark-soil.png', w=260), 3150, GROUND + 50)
place(canvas, cut(D + 's8-05-earthworm.png', w=230, rot=-8), 2700, GROUND + 190, sh=False)
place(canvas, cut(D + 's4-05-fern-fiddlehead.png', w=240), 3700, GROUND + 40)
place(canvas, cut(D + 's4-04-moss.png', w=300), 4000, GROUND + 50)
place(canvas, cut(D + 's4-03-mushrooms.png', w=300), 4800, GROUND + 50)
place(canvas, cut(D + 's8-03-wood-chips.png', w=330, flip=True), 5050, GROUND + 50)
place(canvas, cut(D + 's4-03-mushrooms.png', w=220, flip=True), 5180, GROUND + 40)

# ---------- compost heap, spade, fork ----------
def heap(w, h, seed):
    im = Image.new('RGBA', (w, h), (0, 0, 0, 0)); hdr = ImageDraw.Draw(im)
    r = random.Random(seed)
    pts = [(x, h - h * (math.sin(math.pi * x / w) ** 0.55) + r.uniform(-5, 5)) for x in range(0, w + 1, 8)]
    hdr.polygon(pts, fill=(58, 38, 28, 255))
    mask = im.split()[3]
    tex = Image.new('RGBA', (w, h), (0, 0, 0, 0)); td_ = ImageDraw.Draw(tex)
    for _ in range(w * h // 40):
        x, y = r.randint(0, w), r.randint(0, h); rr_ = r.randint(1, 4)
        c = r.choice([(92, 64, 44), (40, 26, 20), (66, 44, 32), (120, 92, 60), (30, 20, 16), (96, 84, 52)])
        td_.ellipse((x - rr_, y - rr_ * 0.6, x + rr_, y + rr_ * 0.6), fill=c + (255,))
    for _ in range(25):  # straw and twigs
        x, y = r.randint(0, w), r.randint(0, h); a = r.uniform(0, 3.14); l = r.randint(14, 40)
        td_.line((x, y, x + l * math.cos(a), y + l * math.sin(a)), fill=(196, 168, 110, 255), width=3)
    for _ in range(60):  # white fungal threads on the surface
        x, y = r.randint(0, w), r.randint(int(h * 0.2), h); pts2 = [(x, y)]
        for k in range(6): x += r.randint(-10, 10); y += r.randint(-6, 6); pts2.append((x, y))
        td_.line(pts2, fill=(245, 240, 228, 200), width=1)
    tex.putalpha(Image.composite(tex.split()[3], Image.new('L', (w, h), 0), mask))
    im.alpha_composite(tex)
    shade = Image.new('RGBA', (w, h), (0, 0, 0, 0)); sd = ImageDraw.Draw(shade)
    for i in range(w):  # light from the sun on the left
        sd.line((i, 0, i, h), fill=(0, 0, 0, int(90 * i / w)))
    shade.putalpha(Image.composite(shade.split()[3], Image.new('L', (w, h), 0), mask))
    im.alpha_composite(shade)
    return im


def spade(h):
    w = int(h * 0.32); im = Image.new('RGBA', (w, h), (0, 0, 0, 0)); sd = ImageDraw.Draw(im)
    cx = w // 2
    sd.rectangle((cx - 9, int(h * 0.12), cx + 9, int(h * 0.68)), fill=(150, 104, 62, 255))           # wooden shaft
    sd.line((cx - 3, int(h * 0.13), cx - 3, int(h * 0.66)), fill=(186, 140, 90, 255), width=4)       # highlight
    sd.rounded_rectangle((cx - 40, 0, cx + 40, int(h * 0.12)), radius=16, outline=(120, 82, 48, 255), width=12)  # D grip
    sd.rectangle((cx - 12, int(h * 0.66), cx + 12, int(h * 0.74)), fill=(110, 110, 112, 255))        # metal neck
    blade = [(cx - 55, int(h * 0.73)), (cx + 55, int(h * 0.73)), (cx + 50, int(h * 0.93)), (cx, h - 2), (cx - 50, int(h * 0.93))]
    sd.polygon(blade, fill=(132, 136, 140, 255)); sd.line(blade + [blade[0]], fill=(80, 82, 86, 255), width=4)
    sd.line((cx - 30, int(h * 0.76), cx - 22, int(h * 0.92)), fill=(200, 204, 208, 255), width=5)
    return im


def fork(h):
    w = int(h * 0.3); im = Image.new('RGBA', (w, h), (0, 0, 0, 0)); fd = ImageDraw.Draw(im)
    cx = w // 2
    fd.rectangle((cx - 8, int(h * 0.1), cx + 8, int(h * 0.66)), fill=(150, 104, 62, 255))
    fd.rounded_rectangle((cx - 36, 0, cx + 36, int(h * 0.1)), radius=14, outline=(120, 82, 48, 255), width=10)
    fd.rectangle((cx - 44, int(h * 0.66), cx + 44, int(h * 0.71)), fill=(96, 98, 102, 255))
    for k in range(4):
        x = cx - 40 + k * 26
        fd.line((x, int(h * 0.7), x + 2, h - 4), fill=(110, 112, 116, 255), width=8)
    return im


hp = heap(620, 300, 5)
place(canvas, spade(480).rotate(-16, expand=True, resample=Image.BICUBIC), 2700, GROUND - 90)
place(canvas, hp, 2520, GROUND + 30)
steam = Image.new('RGBA', canvas.size, (0, 0, 0, 0)); st = ImageDraw.Draw(steam)
for k, x0 in enumerate((2380, 2470, 2560)):  # warm compost steaming
    pts = [(x0 + 18 * math.sin(t / 22 + k), GROUND - 250 - t) for t in range(0, 150, 6)]
    st.line(pts, fill=(255, 250, 240, 170), width=12)
canvas.alpha_composite(steam.filter(ImageFilter.GaussianBlur(8)))
place(canvas, cut(D + 's1-08-seedling.png', w=120), 2400, GROUND - 255, sh=False)
place(canvas, fork(400).rotate(10, expand=True, resample=Image.BICUBIC), 4130, GROUND + 10)
# flowers in the foreground, overlapping the ground line
for f, x, h, rot in [('norfolk-common-poppy.png', 110, 440, -6), ('costa-rica-hibiscus.png', 330, 300, 4), ('tamil-nadu-marigold.png', 560, 420, -3),
                     ('santa-fe-indian-paintbrush.png', 800, 400, 5), ('norfolk-foxglove.png', 1000, 480, -4),
                     ('norfolk-cornflower.png', 2120, 470, 6), ('norfolk-sea-lavender.png', 3330, 430, -5),
                     ('tamil-nadu-jasmine.png', 4420, 230, 0), ('norfolk-cornflower.png', 5440, 420, -6),
                     ('norfolk-common-poppy.png', 5700, 420, 6), ('tamil-nadu-marigold.png', 5900, 400, -4),
                     ('norfolk-foxglove.png', 6280, 480, 4), ('santa-fe-indian-paintbrush.png', 6100, 430, -6)]:
    place(canvas, cut(D + f, h=h, rot=rot), x, GROUND + 70)

# soft vignette at top for text
vig = Image.new('RGBA', canvas.size, (0, 0, 0, 0)); vg = ImageDraw.Draw(vig)
for i in range(300):
    vg.line((0, i, PW, i), fill=(30, 20, 40, int(110 * (1 - i / 300) ** 2)))
canvas.alpha_composite(vig)
canvas.convert('RGB').save(OUT + 'panorama.jpg', quality=90)

# ---------- text per slide ----------
TEXT = [
    ('What can you do for your soil right now?', 'Five things that make a difference.'),
    ('Cover your soil.', 'Mulch, cover crops, or leaf litter left where it falls.'),
    ('Add biologically active compost.', 'Bacteria, fungi, protozoa and nematodes, all working together.'),
    ('Reduce tillage where you can.', 'Fungal networks that took months to build can be gone in a single pass.'),
    ('Feed the fungi.', 'Woody mulch and carbon-rich material shift the balance toward fungi.'),
    ('Tell someone what you are learning.', 'Demonstration is still the most powerful thing we have.'),
]


def wrap(dr, s, f, width):
    lines, cur = [], ''
    for w_ in s.split():
        t = (cur + ' ' + w_).strip()
        if dr.textlength(t, font=f) <= width: cur = t
        else: lines.append(cur); cur = w_
    return lines + [cur]


logo = cut(D + 's9-02-logo.png', w=120)
PX = 9525
prs = Presentation(); prs.slide_width, prs.slide_height = Emu(W * PX), Emu(H * PX)
for i in range(N):
    bg = canvas.crop((i * W, 0, (i + 1) * W, H))
    bgp = OUT + f'layers/s{i + 1}-background.jpg'; bg.convert('RGB').save(bgp, quality=92)
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    sl.shapes.add_picture(bgp, 0, 0, Emu(W * PX), Emu(H * PX))
    img = bg.copy()
    tl = Image.new('RGBA', img.size, (0, 0, 0, 0)); td = ImageDraw.Draw(tl)
    head, sub = TEXT[i]
    hs = 92
    y = 90
    boxes = []
    for txt, f, size, gap in ((head, FI(hs), hs, 18), (sub, FR(44), 44, 0)):
        y0 = y
        for ln in wrap(td, txt, f, 900):
            td.text((90, y), ln, font=f, fill=CREAM + (255,)); y += int(size * 1.1)
        boxes.append((txt, y0, y - y0, size, txt == head))
        y += gap
    shd = Image.new('RGBA', img.size, (0, 0, 0, 0)); shd.putalpha(tl.split()[3].point(lambda v: int(v * 0.55)))
    img.alpha_composite(shd.filter(ImageFilter.GaussianBlur(6)), (2, 4)); img.alpha_composite(tl)
    if i in (0, N - 1):
        img.alpha_composite(logo, (W - logo.width - 50, H - logo.height - 40))
        lp = OUT + 'layers/logo.png'; logo.save(lp)
        sl.shapes.add_picture(lp, Emu((W - logo.width - 50) * PX), Emu((H - logo.height - 40) * PX), Emu(logo.width * PX), Emu(logo.height * PX))
    for txt, y0, hh, size, is_head in boxes:
        tb = sl.shapes.add_textbox(Emu(90 * PX), Emu(y0 * PX), Emu(900 * PX), Emu(hh * PX))
        tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
        for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
        r = tf.paragraphs[0].add_run(); r.text = txt
        r.font.name = 'EB Garamond'; r.font.italic = is_head; r.font.size = Pt(size * 0.75); r.font.color.rgb = RGBColor(*CREAM)
    img.convert('RGB').save(OUT + f'19-10-2026-improve-your-soil-{i + 1}.jpg', quality=93)
prs.save(OUT + '19-10-2026-improve-your-soil.pptx')
pv = canvas.resize((PW // 5, H // 5)).convert('RGB')
sheet = Image.new('RGB', (PW // 5, H // 5 * 2 + 20), 'white'); sheet.paste(pv, (0, 0))
row = Image.new('RGB', (PW // 5, H // 5), 'white')
for i in range(N):
    row.paste(Image.open(OUT + f'19-10-2026-improve-your-soil-{i + 1}.jpg').resize((W // 5 - 4, H // 5)), (i * (W // 5), 0))
sheet.paste(row, (0, H // 5 + 20)); sheet.save(OUT + 'sheet.jpg'); print('ok')
