"""The dinner party for a billion guests: one long banquet table across six slides."""
import os, math, random
import numpy as np
from PIL import Image, ImageDraw, ImageFont, ImageFilter
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR

D = '/home/claude/p05/fl/'
M = '/home/claude/sm/assets/microscopy/'
OUT = '/home/claude/p05/party/'
os.makedirs(OUT + 'layers', exist_ok=True)
N, W, H = 6, 1080, 1350
PW = W * N
TOP = 930           # table top
FONTS = os.path.expanduser('~/.fonts/')
FI = lambda s: ImageFont.truetype(FONTS + 'eb-garamond-latin-500-italic.ttf', s)
FR = lambda s: ImageFont.truetype(FONTS + 'eb-garamond-latin-400-normal.ttf', s)
FB = lambda s: ImageFont.truetype(FONTS + 'eb-garamond-latin-600-normal.ttf', s)
CREAM = (252, 244, 228)
INK = (52, 34, 30)
random.seed(4)


def grad(h, stops):
    ys = np.linspace(0, 1, h); out = np.zeros((h, 3))
    for c in range(3): out[:, c] = np.interp(ys, [s[0] for s in stops], [s[1][c] for s in stops])
    return out


def shadow(im, off=(10, 14), blur=12, alpha=0.45):
    pad = blur * 3 + max(off)
    out = Image.new('RGBA', (im.width + pad * 2, im.height + pad * 2), (0, 0, 0, 0))
    s = Image.new('RGBA', im.size, (20, 10, 10, 0)); s.putalpha(im.split()[3].point(lambda v: int(v * alpha)))
    sl = Image.new('RGBA', out.size, (0, 0, 0, 0)); sl.alpha_composite(s, (pad + off[0], pad + off[1]))
    out.alpha_composite(sl.filter(ImageFilter.GaussianBlur(blur))); out.alpha_composite(im, (pad, pad))
    return out, pad


def place(img, im, cx, by, sh=True):
    big, pad = shadow(im) if sh else (im, 0)
    x, y = int(cx - im.width / 2 - pad), int(by - im.height - pad)
    img.paste(big, (x, y), big) if (x < 0 or y < 0) else img.alpha_composite(big, (x, y))


def cut(path, w=None, h=None, rot=0, flip=False):
    im = Image.open(path).convert('RGBA')
    if flip: im = im.transpose(Image.FLIP_LEFT_RIGHT)
    if w: im = im.resize((w, max(1, int(im.height * w / im.width))), Image.LANCZOS)
    elif h: im = im.resize((max(1, int(im.width * h / im.height)), h), Image.LANCZOS)
    if rot: im = im.rotate(rot, expand=True, resample=Image.BICUBIC)
    return im


def glow(img, x, y, r, col, a):
    g = Image.new('RGBA', img.size, (0, 0, 0, 0)); gd = ImageDraw.Draw(g)
    gd.ellipse((x - r, y - r, x + r, y + r), fill=col + (a,))
    img.alpha_composite(g.filter(ImageFilter.GaussianBlur(r / 2.2)))


# ---------- the room ----------
wall = grad(H, [(0, (34, 46, 54)), (0.45, (58, 58, 70)), (0.7, (92, 66, 64)), (1, (60, 36, 34))])
C = Image.fromarray(np.repeat(wall[:, None, :], PW, axis=1).astype(np.uint8)).convert('RGBA')
d = ImageDraw.Draw(C)
for x in range(0, PW, 90):  # faint wallpaper stripes and leaf sprigs drawn in gold
    d.line((x, 0, x, TOP - 120), fill=(255, 220, 160, 14), width=2)
    for y in range(60, TOP - 160, 140):
        yy = y + (45 if (x // 90) % 2 else 0)
        d.arc((x + 18, yy, x + 46, yy + 28), 200, 340, fill=(255, 220, 160, 26), width=2)
d.rectangle((0, TOP - 130, PW, TOP - 118), fill=(120, 84, 66, 255))          # dado rail

# windows with a dusk sky and moon
for wx, moon in [(560, True), (3800, False), (5800, False)]:
    ww, wh, wy = 300, 420, 300
    sky = grad(wh, [(0, (40, 46, 90)), (0.6, (196, 120, 120)), (1, (246, 182, 130))])
    S = Image.fromarray(np.repeat(sky[:, None, :], ww, axis=1).astype(np.uint8)).convert('RGBA')
    C.alpha_composite(S, (wx - ww // 2, wy))
    if moon:
        glow(C, wx + 60, wy + 100, 70, (255, 240, 210), 160)
        ImageDraw.Draw(C).ellipse((wx + 30, wy + 70, wx + 90, wy + 130), fill=(255, 246, 226, 255))
    d = ImageDraw.Draw(C)
    d.rectangle((wx - ww // 2 - 14, wy - 14, wx + ww // 2 + 14, wy + wh + 14), outline=(150, 112, 84, 255), width=14)
    d.line((wx, wy, wx, wy + wh), fill=(150, 112, 84, 255), width=10)
    d.line((wx - ww // 2, wy + wh // 2, wx + ww // 2, wy + wh // 2), fill=(150, 112, 84, 255), width=10)
    # curtains
    for side in (-1, 1):
        cx0 = wx + side * (ww // 2 + 40)
        pts = [(cx0 - 40, wy - 50), (cx0 + 40, wy - 50)] + [(cx0 + 40 + 10 * math.sin(t / 30), wy - 50 + t) for t in range(0, wh + 120, 10)] + \
              [(cx0 - 40 + 10 * math.sin(t / 30 + 1), wy - 50 + t) for t in range(wh + 110, -1, -10)]
        d.polygon(pts, fill=(120, 40, 46, 255))
        for k in range(-30, 40, 20):
            d.line([(cx0 + k + 6 * math.sin(t / 30), wy - 50 + t) for t in range(0, wh + 120, 10)], fill=(90, 28, 34, 255), width=3)

# chandelier glow above the middle
for gx in (1620, 3240, 4860):
    glow(C, gx, 120, 220, (255, 210, 150), 70)
    d = ImageDraw.Draw(C)
    d.line((gx, 0, gx, 90), fill=(180, 150, 100, 255), width=4)
    d.arc((gx - 90, 60, gx + 90, 150), 0, 180, fill=(200, 168, 110, 255), width=6)
    for k in (-90, -45, 0, 45, 90):
        d.rectangle((gx + k - 4, 92, gx + k + 4, 112), fill=(250, 240, 220, 255))
        glow(C, gx + k, 86, 16, (255, 220, 150), 220)

# chair backs behind the table
for x in range(160, PW, 340):
    d = ImageDraw.Draw(C)
    d.rounded_rectangle((x - 70, TOP - 250, x + 70, TOP + 10), radius=60, fill=(84, 46, 38, 255))
    d.rounded_rectangle((x - 50, TOP - 230, x + 50, TOP - 40), radius=44, fill=(130, 54, 52, 255))

# ---------- the table ----------
d = ImageDraw.Draw(C)
d.rectangle((0, TOP, PW, H), fill=(80, 46, 36, 255))
cloth = grad(260, [(0, (250, 246, 236)), (1, (226, 216, 200))])
CL = Image.fromarray(np.repeat(cloth[:, None, :], PW, axis=1).astype(np.uint8)).convert('RGBA')
C.alpha_composite(CL, (0, TOP))
d = ImageDraw.Draw(C)
d.line([(x, TOP + 260 + 14 * math.sin(x / 60)) for x in range(0, PW, 8)], fill=(214, 202, 186, 255), width=12)
for x in range(30, PW, 120):  # cloth folds
    d.line((x, TOP + 60, x + 6, TOP + 260), fill=(222, 212, 196, 255), width=3)
d.rectangle((0, TOP + 270, PW, H), fill=(70, 40, 32, 255))
for x in range(0, PW, 12):
    d.line((x, TOP + 270, x, H), fill=(60 + (x // 12) % 3 * 6, 34, 28, 255), width=1)


def plate(cx, cy, r=64):
    d = ImageDraw.Draw(C)
    d.ellipse((cx - r, cy - r * 0.32, cx + r, cy + r * 0.32), fill=(255, 255, 252, 255), outline=(200, 190, 176, 255), width=3)
    d.ellipse((cx - r * 0.62, cy - r * 0.2, cx + r * 0.62, cy + r * 0.2), outline=(214, 180, 120, 255), width=3)


def glass(cx, by):
    d = ImageDraw.Draw(C)
    d.polygon([(cx - 18, by - 90), (cx + 18, by - 90), (cx + 10, by - 50), (cx - 10, by - 50)], fill=(230, 240, 250, 150), outline=(240, 248, 255, 255))
    d.line((cx, by - 50, cx, by - 6), fill=(240, 248, 255, 255), width=4)
    d.ellipse((cx - 16, by - 10, cx + 16, by), outline=(240, 248, 255, 255), width=3)
    d.polygon([(cx - 14, by - 80), (cx + 14, by - 80), (cx + 10, by - 56), (cx - 10, by - 56)], fill=(150, 30, 50, 210))


def candle(cx, by, h=150):
    d = ImageDraw.Draw(C)
    d.ellipse((cx - 34, by - 12, cx + 34, by + 8), fill=(200, 168, 100, 255))
    d.rectangle((cx - 10, by - h, cx + 10, by - 4), fill=(252, 246, 230, 255))
    glow(C, cx, by - h - 22, 60, (255, 200, 120), 150)
    d = ImageDraw.Draw(C)
    d.ellipse((cx - 7, by - h - 36, cx + 7, by - h - 6), fill=(255, 214, 120, 255))
    d.ellipse((cx - 3, by - h - 26, cx + 3, by - h - 10), fill=(255, 250, 230, 255))


def name_card(cx, by, text):
    f = FB(30); tw = int(ImageDraw.Draw(C).textlength(text, font=f)) + 40
    card = Image.new('RGBA', (tw, 70), (0, 0, 0, 0)); cd = ImageDraw.Draw(card)
    cd.polygon([(0, 70), (tw, 70), (tw - 6, 0), (6, 0)], fill=(255, 252, 244, 255), outline=(200, 170, 110, 255))
    cd.text((tw / 2, 38), text, font=f, fill=INK, anchor='mm')
    place(C, card, cx, by, sh=True)


def party_hat(im, cx_frac=0.5, size=None, col=(220, 70, 80)):
    size = size or int(im.width * 0.3)
    hat = Image.new('RGBA', (size, int(size * 1.3)), (0, 0, 0, 0)); hd = ImageDraw.Draw(hat)
    hd.polygon([(0, hat.height), (size, hat.height), (size // 2, 8)], fill=col + (255,))
    for k in range(3):
        y = hat.height - (k + 1) * hat.height // 4
        hd.line((size * (0.5 - (hat.height - y) / hat.height * 0.0) - (hat.height - (hat.height - y)) * 0.0, y, size, y), fill=(0, 0, 0, 0))
    for k in (0.3, 0.55, 0.8):
        y = int(8 + (hat.height - 8) * k); half = size / 2 * k
        hd.line((size / 2 - half, y, size / 2 + half, y), fill=(255, 220, 120, 255), width=5)
    hd.ellipse((size // 2 - 12, 0, size // 2 + 12, 22), fill=(255, 220, 120, 255))
    out = Image.new('RGBA', (im.width, im.height + hat.height - 20), (0, 0, 0, 0))
    out.alpha_composite(im, (0, hat.height - 20))
    out.alpha_composite(hat, (int(im.width * cx_frac - size / 2), 0))
    return out


def guest(path, w, rot=0, hat=(220, 70, 80), flip=False):
    return party_hat(cut(path, w=w, rot=rot, flip=flip), col=hat)


def bacteria(n=7, size=150, seed=1):
    r = random.Random(seed)
    im = Image.new('RGBA', (size, size), (0, 0, 0, 0)); bd = ImageDraw.Draw(im)
    for _ in range(n):
        x, y = r.randint(20, size - 40), r.randint(20, size - 30); a = r.uniform(0, 3.14)
        dx, dy = 18 * math.cos(a), 18 * math.sin(a)
        bd.line((x - dx, y - dy, x + dx, y + dy), fill=(120, 170, 120, 255), width=14)
        bd.line((x - dx, y - dy, x + dx, y + dy), fill=(170, 210, 160, 255), width=6)
    return im


def fungus_guest(h=220):
    im = Image.new('RGBA', (180, h), (0, 0, 0, 0)); fd = ImageDraw.Draw(im)
    x, y = 90, h - 6
    pts = [(x, y)]
    for k in range(40):
        x += 6 * math.sin(k / 3); y -= h / 44; pts.append((x, y))
    fd.line(pts, fill=(250, 246, 236, 255), width=8)
    for k in range(5, 40, 7):
        bx, by = pts[k]
        fd.line((bx, by, bx + (30 if k % 2 else -30), by - 30), fill=(250, 246, 236, 255), width=5)
    return im


# ---------- place settings along the whole table ----------
for x in range(160, PW, 340):
    plate(x, TOP + 120)
    glass(x + 95, TOP + 110)
for x in (420, 1300, 2100, 2950, 4300, 5050, 6000):
    candle(x, TOP + 70)

# slide 2: the huge tablecloth thrown over, guests peeking out
sx = 1080
cloth2 = Image.new('RGBA', (900, 520), (0, 0, 0, 0)); c2 = ImageDraw.Draw(cloth2)
pts = [(0, 520)] + [(x, 160 + 120 * math.sin(x / 900 * math.pi) * -1 + 40 * math.sin(x / 70)) for x in range(0, 901, 10)] + [(900, 520)]
c2.polygon(pts, fill=(240, 232, 214, 255))
for x in range(60, 900, 90):
    c2.line([(x + 20 * math.sin(t / 50), 200 - 80 * math.sin(x / 900 * math.pi) + t) for t in range(0, 320, 10)], fill=(214, 202, 182, 255), width=4)
place(C, cloth2, sx + 560, TOP + 200)
for f, cx, w, col in [('cutout-2.png', sx + 300, 230, (90, 140, 210)), ('cutout-6.png', sx + 560, 250, (220, 70, 80)), ('cutout-3.png', sx + 820, 230, (240, 180, 60))]:
    g = guest(M + f, w, hat=col)
    place(C, g.crop((0, 0, g.width, int(g.height * 0.62))), cx, TOP + 60)

# slide 3: the guests arrive with name cards
sx = 2160
place(C, bacteria(size=220), sx + 170, TOP + 40)
place(C, party_hat(bacteria(seed=3), size=50, col=(90, 140, 210)), sx + 170, TOP + 40, sh=False)
name_card(sx + 170, TOP + 130, 'Bacteria')
place(C, party_hat(fungus_guest(300), size=70, col=(220, 70, 80)), sx + 420, TOP + 50)
name_card(sx + 420, TOP + 130, 'Fungi')
place(C, guest(M + 'cutout-2.png', 280, hat=(240, 180, 60)), sx + 660, TOP + 50)
name_card(sx + 660, TOP + 130, 'Protozoa')
place(C, guest(M + 'cutout-6.png', 290, rot=-10, hat=(120, 90, 200)), sx + 910, TOP + 50)
name_card(sx + 910, TOP + 130, 'Nematode')

# slide 4: the plough crashes through, plates flying
sx = 3240
d = ImageDraw.Draw(C)
for k in range(9):  # crack in the table
    x0 = sx + 540 + random.randint(-60, 60)
    d.line([(x0 + random.randint(-30, 30) * t / 8, TOP + 20 * t) for t in range(0, 16)], fill=(90, 60, 44, 255), width=4)
blade = Image.new('RGBA', (520, 560), (0, 0, 0, 0)); bd = ImageDraw.Draw(blade)
bd.polygon([(60, 40), (300, 0), (480, 360), (260, 540), (120, 380)], fill=(120, 124, 130, 255), outline=(60, 62, 68, 255))
bd.line((300, 0, 260, 540), fill=(170, 174, 180, 255), width=10)
bd.rectangle((280, 0, 320, 40), fill=(90, 60, 40, 255))
place(C, blade.rotate(-12, expand=True), sx + 560, TOP + 160)
for i, (px, py, rot) in enumerate([(sx + 230, 640, 30), (sx + 860, 560, -40), (sx + 940, 760, 20), (sx + 160, 820, -20)]):
    p = Image.new('RGBA', (150, 60), (0, 0, 0, 0)); pd = ImageDraw.Draw(p)
    pd.ellipse((0, 0, 150, 60), fill=(255, 255, 252, 255), outline=(200, 190, 176, 255), width=3)
    pd.ellipse((30, 12, 120, 48), outline=(214, 180, 120, 255), width=3)
    place(C, p.rotate(rot, expand=True), px, py)
    d = ImageDraw.Draw(C)
    for k in range(3):
        d.line((px - 90 - k * 10, py - 30 + k * 18, px - 140 - k * 10, py - 20 + k * 18), fill=(255, 240, 220, 160), width=4)
for f, x, w, rot in [('cutout-3.png', sx + 300, 120, 40), ('cutout-2.png', sx + 820, 120, -50)]:
    place(C, guest(M + f, w, rot=rot), x, TOP - 260)

# slide 5: the wood-chip cake
sx = 4320
cake_x = sx + 560
place(C, cut(D + 's1-03-cake-plate.png', w=520), cake_x, TOP + 120)
tiers = [('s1-05-cake-tier-wood-chips.png', 440), ('s8-03-wood-chips.png', 340), ('s1-05-cake-tier-wood-chips.png', 250)]
by = TOP + 80
for f, w in tiers:
    t = cut(D + f, w=w)
    if f.startswith('s8'): t = t.resize((w, int(w * 0.45)))
    place(C, t, cake_x, by); by -= int(t.height * 0.82)
place(C, fungus_guest(150), cake_x + 10, by + 30, sh=False)
place(C, guest(M + 'cutout-6.png', 160, rot=8, hat=(90, 140, 210)), sx + 230, TOP + 60)
place(C, party_hat(fungus_guest(180), size=50), sx + 900, TOP + 60)

# slide 6: the postman cycles off with the invitations
sx = 5400
d = ImageDraw.Draw(C)
bx, byy = sx + 420, 860
for wx in (bx - 150, bx + 150):  # bicycle, drawn
    d.ellipse((wx - 95, byy - 95, wx + 95, byy + 95), outline=(40, 30, 30, 255), width=8)
    for a in range(0, 360, 30):
        d.line((wx, byy, wx + 90 * math.cos(math.radians(a)), byy + 90 * math.sin(math.radians(a))), fill=(90, 80, 80, 255), width=2)
d.line((bx - 150, byy, bx - 20, byy - 150, bx + 120, byy - 150, bx + 150, byy), fill=(150, 40, 46, 255), width=10)
d.line((bx - 20, byy - 150, bx, byy), fill=(150, 40, 46, 255), width=10)
d.line((bx + 120, byy - 150, bx + 140, byy - 210), fill=(40, 30, 30, 255), width=8)
d.line((bx + 110, byy - 214, bx + 175, byy - 206), fill=(40, 30, 30, 255), width=8)
d.ellipse((bx - 55, byy - 175, bx + 15, byy - 150), fill=(40, 30, 30, 255))
rider = guest(M + 'cutout-6.png', 240, rot=-60, hat=(220, 70, 80))
place(C, rider, bx - 10, byy - 150, sh=False)
d = ImageDraw.Draw(C)
d.rounded_rectangle((bx - 200, byy - 150, bx - 90, byy - 70), radius=10, fill=(120, 80, 50, 255))
for i, (ex, ey, rot) in enumerate([(sx + 720, 640, 18), (sx + 840, 560, -12), (sx + 960, 480, 24), (sx + 620, 540, -20), (sx + 980, 680, 8)]):
    e = Image.new('RGBA', (130, 86), (0, 0, 0, 0)); ed = ImageDraw.Draw(e)
    ed.rectangle((0, 0, 129, 85), fill=(255, 250, 238, 255), outline=(200, 170, 110, 255), width=3)
    ed.line((0, 0, 65, 48, 129, 0), fill=(200, 170, 110, 255), width=3)
    ed.ellipse((54, 38, 76, 60), fill=(160, 30, 40, 255))
    place(C, e.rotate(rot, expand=True), ex, ey)
for k in range(5):
    d.line((bx - 300 - k * 20, byy - 60 + k * 25, bx - 380 - k * 20, byy - 60 + k * 25), fill=(255, 240, 220, 120), width=4)

# top fade for the words
vig = Image.new('RGBA', C.size, (0, 0, 0, 0)); vg = ImageDraw.Draw(vig)
for i in range(420):
    vg.line((0, i, PW, i), fill=(14, 18, 26, int(150 * (1 - i / 420) ** 1.6)))
C.alpha_composite(vig)

TEXT = [
    ('What can you do for your soil right now?', 'Throw a party for a billion guests.'),
    ('Cover your soil.', 'Mulch, cover crops, or leaf litter left where it falls.'),
    ('Add biologically active compost.', 'Bacteria, fungi, protozoa and nematodes, all working together.'),
    ('Reduce tillage where you can.', 'Fungal networks that took months to build can be gone in a single pass.'),
    ('Feed the fungi.', 'Woody mulch and carbon-rich material shift the balance toward fungi.'),
    ('Tell someone what you are learning.', 'Demonstration is still the most powerful thing we have.'),
]


def wrap(dr, s, f, width):
    lines, cur = [], ''
    for w_ in s.split():
        t = (cur + ' ' + w_).strip()
        if dr.textlength(t, font=f) <= width: cur = t
        else: lines.append(cur); cur = w_
    return lines + [cur]


logo = cut(D + 's9-02-logo.png', w=120)
PX = 9525
prs = Presentation(); prs.slide_width, prs.slide_height = Emu(W * PX), Emu(H * PX)
for i in range(N):
    bg = C.crop((i * W, 0, (i + 1) * W, H)); bgp = OUT + f'layers/s{i + 1}-background.jpg'; bg.convert('RGB').save(bgp, quality=92)
    sl = prs.slides.add_slide(prs.slide_layouts[6]); sl.shapes.add_picture(bgp, 0, 0, Emu(W * PX), Emu(H * PX))
    img = bg.copy(); tl = Image.new('RGBA', img.size, (0, 0, 0, 0)); td = ImageDraw.Draw(tl)
    y = 80; boxes = []
    for txt, f, size, gap, head in ((TEXT[i][0], FI(90), 90, 14, True), (TEXT[i][1], FR(44), 44, 0, False)):
        y0 = y
        for ln in wrap(td, txt, f, 900):
            td.text((90, y), ln, font=f, fill=CREAM + (255,)); y += int(size * 1.1)
        boxes.append((txt, y0, y - y0, size, head)); y += gap
    shd = Image.new('RGBA', img.size, (0, 0, 0, 0)); shd.putalpha(tl.split()[3].point(lambda v: int(v * 0.6)))
    img.alpha_composite(shd.filter(ImageFilter.GaussianBlur(6)), (2, 4)); img.alpha_composite(tl)
    if i in (0, N - 1):
        img.alpha_composite(logo, (W - logo.width - 50, H - logo.height - 40)); lp = OUT + 'layers/logo.png'; logo.save(lp)
        sl.shapes.add_picture(lp, Emu((W - logo.width - 50) * PX), Emu((H - logo.height - 40) * PX), Emu(logo.width * PX), Emu(logo.height * PX))
    for txt, y0, hh, size, head in boxes:
        tb = sl.shapes.add_textbox(Emu(90 * PX), Emu(y0 * PX), Emu(900 * PX), Emu(hh * PX))
        tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
        for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
        r = tf.paragraphs[0].add_run(); r.text = txt
        r.font.name = 'EB Garamond'; r.font.italic = head; r.font.size = Pt(size * 0.75); r.font.color.rgb = RGBColor(*CREAM)
    img.convert('RGB').save(OUT + f'party-{i + 1}.jpg', quality=93)
prs.save(OUT + '19-10-2026-dinner-party.pptx')
row = Image.new('RGB', (N * 220, 275), 'white')
for i in range(N): row.paste(Image.open(OUT + f'party-{i + 1}.jpg').resize((216, 270)), (i * 220, 2))
row.save(OUT + 'sheet.jpg'); print('ok')
