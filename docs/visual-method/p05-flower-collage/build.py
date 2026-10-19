import os, math
from PIL import Image, ImageDraw, ImageFont, ImageFilter
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

D = '/home/claude/p05/fl/'
OUT = '/home/claude/p05/v2/'
os.makedirs(OUT + 'layers', exist_ok=True)
W, H = 1080, 1350
FONTS = os.path.expanduser('~/.fonts/')
FI = lambda s: ImageFont.truetype(FONTS + 'eb-garamond-latin-500-italic.ttf', s)
CREAM = (244, 241, 234)
DG = (34, 55, 31)
FG = (21, 104, 38)
LOGO = 's9-02-logo.png'


def shadow(im, off=(8, 12), blur=10, alpha=0.38):
    pad = blur * 3 + max(off)
    out = Image.new('RGBA', (im.width + pad * 2, im.height + pad * 2), (0, 0, 0, 0))
    a = im.split()[3].point(lambda v: int(v * alpha))
    s = Image.new('RGBA', im.size, (20, 18, 10, 0)); s.putalpha(a)
    sl = Image.new('RGBA', out.size, (0, 0, 0, 0)); sl.alpha_composite(s, (pad + off[0], pad + off[1]))
    out.alpha_composite(sl.filter(ImageFilter.GaussianBlur(blur)))
    out.alpha_composite(im, (pad, pad))
    return out, pad


def cut(name, w=None, h=None, rot=0, flip=False):
    im = Image.open(D + name).convert('RGBA')
    if flip: im = im.transpose(Image.FLIP_LEFT_RIGHT)
    if w: im = im.resize((w, int(im.height * w / im.width)), Image.LANCZOS)
    elif h: im = im.resize((int(im.width * h / im.height), h), Image.LANCZOS)
    if rot: im = im.rotate(rot, expand=True, resample=Image.BICUBIC)
    return im


class Slide:
    def __init__(self, n, bg):
        self.n = n
        self.c = Image.open(D + bg).convert('RGBA').resize((W, H))
        self.layers = [('background', D + bg, 0, 0, W, H)]
        self.texts = []
        self.d = ImageDraw.Draw(self.c)

    def put(self, label, im, x, y, shadowed=True):
        """x, y = top-left of the cutout itself."""
        if shadowed:
            big, pad = shadow(im); x0, y0 = x - pad, y - pad; im2 = big
        else:
            im2, x0, y0 = im, x, y
        self.c.alpha_composite(im2, (int(x0), int(y0))) if x0 >= 0 and y0 >= 0 else self.c.paste(im2, (int(x0), int(y0)), im2)
        p = OUT + f'layers/s{self.n}-{len(self.layers):02d}-{label}.png'
        im2.save(p)
        self.layers.append((label, p, x0, y0, im2.width, im2.height))

    def text(self, s, x, y, size, color, width=900, lead=1.12):
        f = FI(size)
        lines = []
        for para in s.split('\n'):
            cur = ''
            for word in para.split():
                t = (cur + ' ' + word).strip()
                if self.d.textlength(t, font=f) <= width: cur = t
                else: lines.append(cur); cur = word
            lines.append(cur)
        yy = y
        for ln in lines:
            self.d.text((x, yy), ln, font=f, fill=color)
            yy += int(size * lead)
        self.texts.append((s, x, y, width, yy - y, size, color))
        return yy

    def save(self):
        self.c.convert('RGB').save(OUT + f'19-10-2026-improve-your-soil-{self.n}.jpg', quality=93)


slides = []

# 1 cover, green
s = Slide(1, 'bg1.jpg')
y = s.text('What can you do for your soil right now?', 80, 150, 112, CREAM, width=860)
s.text('Five things that make a difference.', 84, y + 30, 46, CREAM)
s.put('logo', cut(LOGO, w=120), 900, 50, shadowed=False)
s.put('foxglove', cut('norfolk-foxglove.png', h=520, rot=-4), 900, 820)
s.put('paintbrush', cut('santa-fe-indian-paintbrush.png', h=430, rot=6), 640, 900)
s.put('poppy', cut('norfolk-common-poppy.png', h=470, rot=-8), 40, 860)
s.put('hibiscus', cut('costa-rica-hibiscus.png', w=300, rot=5), 220, 1010)
s.put('marigold', cut('tamil-nadu-marigold.png', h=380, rot=-3), 500, 990)
slides.append(s)

# 2 cover your soil, cream
s = Slide(2, 'bg2.jpg')
y = s.text('Cover your soil.', 80, 170, 120, FG)
s.text('Mulch, cover crops, or leaf litter left where it falls.', 84, y + 30, 50, DG, width=820)
s.put('cornflower', cut('norfolk-cornflower.png', h=470, rot=8), 860, 40)
s.put('dry-leaves', cut('s8-06-dry-leaves.png', w=430, rot=-6), 40, 930)
s.put('wood-chips', cut('s8-03-wood-chips.png', w=380, rot=4), 400, 990)
s.put('lavender', cut('norfolk-sea-lavender.png', h=430, rot=-4), 790, 880)
slides.append(s)

# 3 compost, green
s = Slide(3, 'bg1.jpg')
y = s.text('Add biologically active compost.', 80, 130, 112, CREAM, width=860)
s.text('Bacteria, fungi, protozoa and nematodes, all working together.', 84, y + 30, 48, CREAM, width=820)
s.put('marigold', cut('tamil-nadu-marigold.png', h=380, rot=10), 880, 40)
s.put('compost', cut('s8-02-compost-fungal-threads.png', w=380, rot=-3), 40, 960)
s.put('dark-soil', cut('s8-04-dark-soil.png', w=360, rot=3), 360, 1000)
s.put('earthworm', cut('s8-05-earthworm.png', w=330, rot=-6), 680, 1010)
slides.append(s)

# 4 reduce tillage, cream
s = Slide(4, 'bg2.jpg')
y = s.text('Reduce tillage where you can.', 80, 150, 112, FG, width=820)
s.text('Fungal networks that took months to build can be gone in a single pass.', 84, y + 30, 48, DG, width=820)
s.put('poppy', cut('norfolk-common-poppy.png', h=460, rot=10), 880, 20)
s.put('fiddlehead', cut('s4-05-fern-fiddlehead.png', w=300, rot=-4), 50, 1000)
s.put('moss', cut('s4-04-moss.png', w=360, rot=2), 350, 1050)
s.put('mushrooms', cut('s4-03-mushrooms.png', w=330, rot=-3), 700, 980)
slides.append(s)

# 5 feed the fungi, green
s = Slide(5, 'bg1.jpg')
y = s.text('Feed the fungi.', 80, 170, 120, CREAM)
s.text('Woody mulch and carbon-rich material shift the balance toward fungi.', 84, y + 30, 48, CREAM, width=820)
s.put('orchid', cut('13-orchid.png', w=260, rot=-10), 820, 60)
s.put('wood-chips', cut('s8-03-wood-chips.png', w=360, rot=-4), 40, 1010)
s.put('mushrooms', cut('s4-03-mushrooms.png', w=360, rot=4), 380, 950)
s.put('moss', cut('s4-04-moss.png', w=330, rot=-2), 720, 1060)
slides.append(s)

# 6 tell someone, cream
s = Slide(6, 'bg2.jpg')
y = s.text('Tell someone what you are learning.', 80, 150, 112, FG, width=860)
s.text('Demonstration is still the most powerful thing we have.', 84, y + 30, 48, DG, width=820)
s.put('logo', cut(LOGO, w=120), 900, 50, shadowed=False)
s.put('poppy', cut('norfolk-common-poppy.png', h=440, rot=-8), 30, 880)
s.put('hibiscus', cut('costa-rica-hibiscus.png', w=290, rot=-5), 230, 1020)
s.put('marigold', cut('tamil-nadu-marigold.png', h=380, rot=4), 500, 960)
s.put('paintbrush', cut('santa-fe-indian-paintbrush.png', h=420, rot=-6), 690, 910)
s.put('foxglove', cut('norfolk-foxglove.png', h=500, rot=5), 880, 830)
slides.append(s)

# render + editable pptx
PX = 9525
prs = Presentation(); prs.slide_width, prs.slide_height = Emu(W * PX), Emu(H * PX)
for s in slides:
    s.save()
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    for label, p, x, y, w, h in s.layers:
        sl.shapes.add_picture(p, Emu(int(x) * PX), Emu(int(y) * PX), Emu(int(w) * PX), Emu(int(h) * PX))
    for t, x, y, w, h, size, color in s.texts:
        tb = sl.shapes.add_textbox(Emu(x * PX), Emu(y * PX), Emu(w * PX), Emu(h * PX))
        tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
        for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
        p = tf.paragraphs[0]; p.line_spacing = 1.0
        r = p.add_run(); r.text = t
        r.font.name = 'EB Garamond'; r.font.italic = True; r.font.size = Pt(size * 0.75); r.font.color.rgb = RGBColor(*color)
prs.save(OUT + '19-10-2026-improve-your-soil.pptx')
ims = [Image.open(OUT + f'19-10-2026-improve-your-soil-{i}.jpg').resize((360, 450)) for i in range(1, 7)]
sh = Image.new('RGB', (3 * 370 + 10, 2 * 460 + 10), '#222')
for i, im in enumerate(ims): sh.paste(im, (10 + (i % 3) * 370, 10 + (i // 3) * 460))
sh.save(OUT + 'sheet.jpg'); print('ok')
