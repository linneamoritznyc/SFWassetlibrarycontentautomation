"""Shared drawing helpers for SFW carousels: scenes, cutouts with soft shadows, text and editable export."""
import os, math, random
import numpy as np
from PIL import Image, ImageDraw, ImageFont, ImageFilter
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR

FL = '/home/claude/p05/fl/'
MIC = '/home/claude/sm/assets/microscopy/'
SRC = '/home/claude/p05/src/'
W, H = 1080, 1350
FONTS = os.path.expanduser('~/.fonts/')
FI = lambda s: ImageFont.truetype(FONTS + 'eb-garamond-latin-500-italic.ttf', s)
FR = lambda s: ImageFont.truetype(FONTS + 'eb-garamond-latin-400-normal.ttf', s)
FB = lambda s: ImageFont.truetype(FONTS + 'eb-garamond-latin-600-normal.ttf', s)
CREAM = (252, 244, 228)
DG = (34, 55, 31)
FG = (21, 104, 38)
INK = (40, 30, 26)
LOGO = FL + 's9-02-logo.png'


def grad(h, stops):
    ys = np.linspace(0, 1, h); out = np.zeros((h, 3))
    for c in range(3): out[:, c] = np.interp(ys, [s[0] for s in stops], [s[1][c] for s in stops])
    return out


def vgrad(w, h, stops):
    g = grad(h, stops)
    return Image.fromarray(np.repeat(g[:, None, :], w, axis=1).astype(np.uint8)).convert('RGBA')


def hgrad(w, h, stops):
    g = grad(w, stops)
    return Image.fromarray(np.repeat(g[None, :, :], h, axis=0).astype(np.uint8)).convert('RGBA')


def grain(img, amt=10, seed=1):
    rng = np.random.default_rng(seed)
    a = np.array(img.convert('RGB'), np.float32)
    n = rng.normal(0, amt, a.shape[:2])[:, :, None]
    out = Image.fromarray(np.clip(a + n, 0, 255).astype(np.uint8)).convert('RGBA')
    return out


def shadow(im, off=(10, 14), blur=12, alpha=0.42):
    pad = blur * 3 + max(off)
    out = Image.new('RGBA', (im.width + pad * 2, im.height + pad * 2), (0, 0, 0, 0))
    s = Image.new('RGBA', im.size, (20, 12, 8, 0)); s.putalpha(im.split()[3].point(lambda v: int(v * alpha)))
    sl = Image.new('RGBA', out.size, (0, 0, 0, 0)); sl.alpha_composite(s, (pad + off[0], pad + off[1]))
    out.alpha_composite(sl.filter(ImageFilter.GaussianBlur(blur))); out.alpha_composite(im, (pad, pad))
    return out, pad


def place(img, im, cx, by, sh=True, **kw):
    """cx = horizontal centre, by = bottom edge of the cutout."""
    big, pad = shadow(im, **kw) if sh else (im, 0)
    x, y = int(cx - im.width / 2 - pad), int(by - im.height - pad)
    if x < 0 or y < 0 or x + big.width > img.width or y + big.height > img.height:
        tmp = Image.new('RGBA', img.size, (0, 0, 0, 0)); tmp.paste(big, (x, y), big); img.alpha_composite(tmp)
    else:
        img.alpha_composite(big, (x, y))


def cut(path, w=None, h=None, rot=0, flip=False):
    im = Image.open(path).convert('RGBA')
    if flip: im = im.transpose(Image.FLIP_LEFT_RIGHT)
    if w: im = im.resize((w, max(1, int(im.height * w / im.width))), Image.LANCZOS)
    elif h: im = im.resize((max(1, int(im.width * h / im.height)), h), Image.LANCZOS)
    if rot: im = im.rotate(rot, expand=True, resample=Image.BICUBIC)
    return im


def glow(img, x, y, r, col, a):
    g = Image.new('RGBA', img.size, (0, 0, 0, 0)); ImageDraw.Draw(g).ellipse((x - r, y - r, x + r, y + r), fill=col + (a,))
    img.alpha_composite(g.filter(ImageFilter.GaussianBlur(max(2, r / 2.2))))


def ridge(base, amp, freq, seed, w):
    rng = np.random.default_rng(seed); x = np.arange(w); y = np.full(w, float(base))
    for k in range(1, 5):
        y += amp / k * np.sin(x / w * 2 * math.pi * freq * k + rng.uniform(0, 6.28))
    return y


def fill_under(img, ys, color, alpha=255, blur=0, bottom=None):
    lay = Image.new('RGBA', img.size, (0, 0, 0, 0)); d = ImageDraw.Draw(lay)
    b = bottom if bottom is not None else img.height
    pts = [(i, ys[i]) for i in range(0, len(ys), 4)] + [(len(ys) - 1, ys[-1]), (len(ys) - 1, b), (0, b)]
    d.polygon(pts, fill=color + (alpha,))
    if blur: lay = lay.filter(ImageFilter.GaussianBlur(blur))
    img.alpha_composite(lay)


def tree(img, x, base, h, seed, tone=(48, 52, 40), leaf=(70, 86, 52), leaf2=None):
    r = random.Random(seed); lay = Image.new('RGBA', img.size, (0, 0, 0, 0)); t = ImageDraw.Draw(lay)

    def br(x0, y0, ang, ln, wd, depth):
        x1 = x0 + ln * math.cos(ang); y1 = y0 - ln * math.sin(ang)
        t.line((x0, y0, x1, y1), fill=tone + (255,), width=max(1, int(wd)))
        if depth == 0 or ln < 14:
            for _ in range(3):
                rr = r.uniform(18, 46) * h / 520; ox, oy = x1 + r.uniform(-22, 22), y1 + r.uniform(-22, 22)
                c = leaf if (leaf2 is None or r.random() < 0.5) else leaf2
                t.ellipse((ox - rr, oy - rr, ox + rr, oy + rr), fill=c + (r.randint(150, 225),))
            return
        for da in (r.uniform(0.25, 0.55), -r.uniform(0.25, 0.55)) + ((r.uniform(-0.15, 0.15),) if r.random() < 0.5 else ()):
            br(x1, y1, ang + da, ln * r.uniform(0.68, 0.8), wd * 0.68, depth - 1)

    br(x, base, math.pi / 2 + r.uniform(-0.05, 0.05), h * 0.32, h / 26, 7)
    img.alpha_composite(lay.filter(ImageFilter.GaussianBlur(0.8)))


def conifer(img, x, base, h, col=(40, 62, 50), alpha=255):
    d = ImageDraw.Draw(img); w = h * 0.36
    for k in range(6):
        y0 = base - h + k * h / 7; ww = w * (k + 1) / 6
        d.polygon([(x, y0), (x - ww / 2, y0 + h / 4), (x + ww / 2, y0 + h / 4)], fill=col + (alpha,))
    d.rectangle((x - h / 40, base - h / 6, x + h / 40, base), fill=(60, 44, 34, alpha))


def lens(path, d, ring=10, ring_col=(250, 244, 230), fx=0.5, fy=0.5):
    im = Image.open(path).convert('RGB'); w, h = im.size; s = min(w, h)
    x0 = int((w - s) * fx); y0 = int((h - s) * fy)
    im = im.crop((x0, y0, x0 + s, y0 + s)).resize((d, d), Image.LANCZOS).convert('RGBA')
    m = Image.new('L', (d, d), 0); ImageDraw.Draw(m).ellipse((0, 0, d - 1, d - 1), fill=255); im.putalpha(m)
    out = Image.new('RGBA', (d + ring * 2, d + ring * 2), (0, 0, 0, 0))
    ImageDraw.Draw(out).ellipse((0, 0, d + ring * 2 - 1, d + ring * 2 - 1), fill=ring_col + (255,))
    out.alpha_composite(im, (ring, ring)); return out


def leafshape(size, col, seed, kind='maple'):
    r = random.Random(seed); im = Image.new('RGBA', (size, size), (0, 0, 0, 0)); d = ImageDraw.Draw(im); c = size / 2
    if kind == 'maple':
        pts = []
        for k in range(10):
            a = math.pi / 2 + k * 2 * math.pi / 10
            rad = size * (0.48 if k % 2 == 0 else 0.22) * r.uniform(0.85, 1.05)
            pts.append((c + rad * math.cos(a), c - rad * math.sin(a)))
        d.polygon(pts, fill=col + (255,))
    else:
        d.ellipse((size * 0.2, size * 0.05, size * 0.8, size * 0.95), fill=col + (255,))
    dark = tuple(max(0, v - 50) for v in col)
    d.line((c, size * 0.95, c, size * 0.08), fill=dark + (255,), width=max(2, size // 40))
    for k in range(3):
        y = size * (0.3 + k * 0.18); d.line((c, y + size * 0.1, c - size * 0.18, y), fill=dark + (255,), width=max(1, size // 60)); d.line((c, y + size * 0.1, c + size * 0.18, y), fill=dark + (255,), width=max(1, size // 60))
    return im


def soil_band(img, top, seed=2, y1=None):
    r = random.Random(seed); y1 = y1 or img.height
    s = vgrad(img.width, y1 - top, [(0, (82, 56, 42)), (0.35, (62, 40, 32)), (1, (32, 22, 20))])
    img.alpha_composite(grain(s, 12, seed), (0, top)); d = ImageDraw.Draw(img)
    for _ in range(img.width * (y1 - top) // 2500):
        x = r.randint(0, img.width); y = r.randint(top + 10, y1); rr = r.randint(2, 8)
        d.ellipse((x - rr, y - rr, x + rr, y + rr * 0.7), fill=r.choice([(110, 84, 66), (90, 70, 56), (140, 118, 96), (54, 38, 32)]) + (255,))


def hyphae(img, box, n, seed, col=(246, 238, 220)):
    r = random.Random(seed); lay = Image.new('RGBA', img.size, (0, 0, 0, 0)); d = ImageDraw.Draw(lay)
    for _ in range(n):
        x, y = r.randint(box[0], box[2]), r.randint(box[1], box[3]); a = r.uniform(0, 6.28); pts = [(x, y)]
        for _ in range(r.randint(8, 20)):
            a += r.uniform(-0.5, 0.5); x += 16 * math.cos(a); y += 16 * math.sin(a) * 0.6; pts.append((x, y))
        d.line(pts, fill=col + (r.randint(90, 180),), width=2)
    img.alpha_composite(lay)


def wrap(dr, s, f, width):
    lines = []
    for para in s.split('\n'):
        cur = ''
        for w_ in para.split():
            t = (cur + ' ' + w_).strip()
            if dr.textlength(t, font=f) <= width: cur = t
            else: lines.append(cur); cur = w_
        lines.append(cur)
    return lines


def top_fade(img, h=420, col=(14, 18, 26), a=150):
    v = Image.new('RGBA', img.size, (0, 0, 0, 0)); vd = ImageDraw.Draw(v)
    for i in range(h): vd.line((0, i, img.width, i), fill=col + (int(a * (1 - i / h) ** 1.6),))
    img.alpha_composite(v)


def _mont(size, w):
    f = ImageFont.truetype(FONTS + 'montserrat-latin-variable.ttf', size); f.set_variation_by_axes([w]); return f


def FONTSEL(k, size):
    return {'i': FI, 'b': FB, 'r': FR, 'g': FI,
            'm': lambda s: _mont(s, 700), 'e': lambda s: _mont(s, 600),
            's': lambda s: ImageFont.truetype(FONTS + 'source-sans-3-latin-400-normal.ttf', s)}[k](size)


def export(name, slides, outdir, w=W, h=H):
    """slides: list of dicts {bg: RGBA image, texts: [(text, x, y, width, size, color, italic, align)], logo: (x, y) or None}"""
    os.makedirs(outdir + '/layers', exist_ok=True)
    PX = 9525; prs = Presentation(); prs.slide_width, prs.slide_height = Emu(w * PX), Emu(h * PX)
    paths = []
    for i, s in enumerate(slides, 1):
        bg = s['bg']; bgp = f'{outdir}/layers/{name}-{i}-background.jpg'; bg.convert('RGB').save(bgp, quality=92)
        sl = prs.slides.add_slide(prs.slide_layouts[6]); sl.shapes.add_picture(bgp, 0, 0, Emu(w * PX), Emu(h * PX))
        img = bg.copy(); tl = Image.new('RGBA', img.size, (0, 0, 0, 0)); td = ImageDraw.Draw(tl)
        for t in s.get('texts', []):
            txt, x, y, width, size, col, ital = t[:7]; align = t[7] if len(t) > 7 else 'l'
            f = FONTSEL(ital, size)
            yy = y
            for ln in wrap(td, txt, f, width):
                lx = x if align == 'l' else x + (width - td.textlength(ln, font=f)) / 2
                td.text((lx, yy), ln, font=f, fill=col + (255,)); yy += int(size * 1.12)
            tb = sl.shapes.add_textbox(Emu(int(x) * PX), Emu(int(y) * PX), Emu(int(width) * PX), Emu(int(yy - y) * PX))
            tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
            for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
            p = tf.paragraphs[0]
            if align == 'c':
                from pptx.enum.text import PP_ALIGN; p.alignment = PP_ALIGN.CENTER
            r = p.add_run(); r.text = txt
            r.font.name = {'m': 'Montserrat', 'e': 'Montserrat', 's': 'Source Sans 3'}.get(ital, 'EB Garamond'); r.font.italic = ital in ('i', 'g'); r.font.bold = ital in ('b', 'm'); r.font.size = Pt(size * 0.75); r.font.color.rgb = RGBColor(*col)
        dark = sum(s.get('texts', [('', 0, 0, 0, 0, CREAM, '')])[0][5]) > 500 if s.get('texts') else False
        if dark and not s.get('flat'):
            shd = Image.new('RGBA', img.size, (0, 0, 0, 0)); shd.putalpha(tl.split()[3].point(lambda v: int(v * 0.6)))
            img.alpha_composite(shd.filter(ImageFilter.GaussianBlur(6)), (2, 4))
        img.alpha_composite(tl)
        if s.get('logo'):
            lg = cut(LOGO, w=120); lx, ly = s['logo']; img.alpha_composite(lg, (lx, ly))
            lp = f'{outdir}/layers/logo.png'; lg.save(lp)
            sl.shapes.add_picture(lp, Emu(lx * PX), Emu(ly * PX), Emu(lg.width * PX), Emu(lg.height * PX))
        p = f'{outdir}/{name}-{i}.jpg'; img.convert('RGB').save(p, quality=92); paths.append(p)
    prs.save(f'{outdir}/{name}.pptx')
    return paths


def sheet(paths, out, tw=216):
    th = int(tw * Image.open(paths[0]).height / Image.open(paths[0]).width)
    row = Image.new('RGB', (len(paths) * (tw + 4), th + 4), 'white')
    for i, p in enumerate(paths): row.paste(Image.open(p).resize((tw, th)), (i * (tw + 4), 2))
    row.save(out)
